import click
import sys
import os
from datetime import datetime
from dotenv import load_dotenv

# Add the project root to Python path
sys.path.append(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from csa.services.java_parser import parse_java_project, extract_beans_from_classes, analyze_bean_dependencies, extract_endpoints_from_classes, extract_mybatis_mappers_from_classes, extract_jpa_entities_from_classes, extract_test_classes_from_classes, extract_sql_statements_from_mappers, extract_project_name
from csa.services.graph_db import GraphDB
from csa.services.neo4j_connection_pool import get_connection_pool, initialize_pool_from_env
from csa.services.sequence_diagram_generator import SequenceDiagramGenerator
from csa.services.db_parser import DBParser
from csa.services.db_call_analysis import DBCallAnalysisService
from csa.utils.logger import get_logger
from csa.models.graph_entities import Project
from neo4j import GraphDatabase
import subprocess
import tempfile

load_dotenv()

def start(command_name):
    """
    CLI 명령어 시작 시 호출되는 공통 초기화 함수
    
    Args:
        command_name: 실행할 명령어 이름 (예: 'analyze', 'crud-matrix')
    
    Returns:
        dict: 시작 시각, connection pool 등 공통 리소스를 담은 컨텍스트
    """
    start_time = datetime.now()
    logger = get_logger(__name__, command=command_name)
    logger.info("")
    logger.info(f"====== {command_name} 작업 시작 ======")
    
    # Connection Pool 초기화 (이미 초기화된 경우 재사용)
    pool = get_connection_pool()
    if not pool.is_initialized():
        neo4j_uri = os.getenv("NEO4J_URI", "bolt://localhost:7687")
        neo4j_user = os.getenv("NEO4J_USER", "neo4j")
        neo4j_password = os.getenv("NEO4J_PASSWORD")
        neo4j_database = os.getenv("NEO4J_DATABASE", "neo4j")
        pool_size = int(os.getenv('NEO4J_POOL_SIZE', '10'))
        
        if neo4j_password:
            pool.initialize(neo4j_uri, neo4j_user, neo4j_password, neo4j_database, pool_size)
            logger.info(f"Connection pool initialized with {pool_size} connections")
        else:
            logger.warning("NEO4J_PASSWORD not set - connection pool not initialized")
    
    return {
        'start_time': start_time,
        'logger': logger,
        'pool': pool,
        'command_name': command_name
    }

def end(context, result=None):
    """
    CLI 명령어 종료 시 호출되는 공통 정리 함수
    
    Args:
        context: start()에서 반환된 컨텍스트 딕셔너리
        result: 명령어 함수에서 반환된 결과 딕셔너리 (optional)
    """
    end_time = datetime.now()
    logger = context.get('logger')
    command_name = context.get('command_name', 'unknown')
    start_time = context.get('start_time')
    
    # 결과 처리 및 통계 출력
    if result:
        if result.get('success'):
            logger.info(f"작업 완료: {result.get('message', 'Success')}")
            
            # 통계 정보 출력
            if 'stats' in result and result['stats']:
                stats = result['stats']
                logger.info("-" * 80)
                logger.info("[작업 통계]")
                for key, value in stats.items():
                    logger.info(f"  • {key}: {value}")
        else:
            logger.error(f"작업 실패: {result.get('error', 'Unknown error')}")
    
    # 작업 시간 계산 및 출력
    if start_time:
        duration = (end_time - start_time).total_seconds()
        logger.info(f"총 수행 시간: {format_duration(duration)}")
    
    logger.info(f"====== {command_name} 작업 완료 ======")
    logger.info("")
    
    # Connection Pool은 애플리케이션 종료 시까지 유지
    # 개별 명령어마다 close하지 않음

def format_duration(seconds):
    """초를 시:분:초 형식으로 변환"""
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    
    if hours > 0:
        return f"{hours}시간 {minutes}분 {secs}초"
    elif minutes > 0:
        return f"{minutes}분 {secs}초"
    else:
        return f"{secs}초"

def format_number(num):
    """숫자를 콤마로 구분하여 포맷팅"""
    return f"{num:,}"

def print_analysis_summary(overall_start_time, overall_end_time, java_stats=None, db_stats=None, dry_run=False):
    """
    분석 작업 완료 후 Summary를 출력합니다.
    
    Args:
        overall_start_time: 전체 작업 시작 시각
        overall_end_time: 전체 작업 종료 시각
        java_stats: Java Object 분석 통계 (dict)
        db_stats: DB Object 분석 통계 (dict)
        dry_run: dry-run 모드 여부
    """
    logger = get_logger(__name__, command='analyze')
    
    # 타이틀 결정
    title = "분석 작업 완료 Summary [dry-run 모드]" if dry_run else "분석 작업 완료 Summary"
    
    # 전체 작업 시간 계산
    total_duration = (overall_end_time - overall_start_time).total_seconds()
    
    logger.info("=" * 80)
    logger.info(f"                          {title}")
    logger.info("=" * 80)
    logger.info("")
    logger.info(f"전체 작업 시간: {overall_start_time.strftime('%Y-%m-%d %H:%M:%S')} ~ {overall_end_time.strftime('%Y-%m-%d %H:%M:%S')}")
    logger.info(f"총 수행 시간: {format_duration(total_duration)}")
    logger.info("")
    
    # Java Object 분석 결과
    if java_stats:
        java_duration = (java_stats['end_time'] - java_stats['start_time']).total_seconds()
        logger.info("-" * 80)
        logger.info("[Java Object 분석 결과]")
        logger.info("-" * 80)
        logger.info(f"작업 시간: {java_stats['start_time'].strftime('%Y-%m-%d %H:%M:%S')} ~ {java_stats['end_time'].strftime('%Y-%m-%d %H:%M:%S')}")
        logger.info(f"수행 시간: {format_duration(java_duration)}")
        logger.info("")
        logger.info("분석 결과:")
        logger.info(f"  • Project: {java_stats.get('project_name', 'N/A')}")
        
        if java_stats.get('total_files', 0) > 0:
            success_rate = (java_stats.get('processed_files', 0) / java_stats.get('total_files', 1)) * 100
            logger.info(f"  • Java 파일: {format_number(java_stats.get('processed_files', 0))}/{format_number(java_stats.get('total_files', 0))}개 (성공률: {success_rate:.1f}%)")
            if java_stats.get('error_files', 0) > 0:
                logger.info(f"  • 에러 파일: {format_number(java_stats.get('error_files', 0))}개")
        
        logger.info(f"  • Packages: {format_number(java_stats.get('packages', 0))}개")
        logger.info(f"  • Classes: {format_number(java_stats.get('classes', 0))}개")
        logger.info(f"  • Methods: {format_number(java_stats.get('methods', 0))}개")
        logger.info(f"  • Fields: {format_number(java_stats.get('fields', 0))}개")
        logger.info(f"  • Spring Beans: {format_number(java_stats.get('beans', 0))}개")
        logger.info(f"  • REST Endpoints: {format_number(java_stats.get('endpoints', 0))}개")
        logger.info(f"  • MyBatis Mappers: {format_number(java_stats.get('mybatis_mappers', 0))}개")
        logger.info(f"  • JPA Entities: {format_number(java_stats.get('jpa_entities', 0))}개")
        logger.info(f"  • JPA Repositories: {format_number(java_stats.get('jpa_repositories', 0))}개")
        logger.info(f"  • SQL Statements: {format_number(java_stats.get('sql_statements', 0))}개")
        logger.info("")
    
    # Database Object 분석 결과
    if db_stats:
        logger.info("-" * 80)
        logger.info("[Database Object 분석 결과]")
        logger.info("-" * 80)
        
        # 시간 정보가 있는 경우에만 표시
        if 'start_time' in db_stats and 'end_time' in db_stats:
            db_duration = (db_stats['end_time'] - db_stats['start_time']).total_seconds()
            logger.info(f"작업 시간: {db_stats['start_time'].strftime('%Y-%m-%d %H:%M:%S')} ~ {db_stats['end_time'].strftime('%Y-%m-%d %H:%M:%S')}")
            logger.info(f"수행 시간: {format_duration(db_duration)}")
            logger.info("")
        
        logger.info("분석 결과:")
        logger.info(f"  • DDL 파일: {format_number(db_stats.get('ddl_files', 0))}개")
        logger.info(f"  • Databases: {format_number(db_stats.get('databases', 0))}개")
        logger.info(f"  • Tables: {format_number(db_stats.get('tables', 0))}개")
        logger.info(f"  • Columns: {format_number(db_stats.get('columns', 0))}개")
        logger.info(f"  • Indexes: {format_number(db_stats.get('indexes', 0))}개")
        logger.info(f"  • Constraints: {format_number(db_stats.get('constraints', 0))}개")
        logger.info("")
    
    logger.info("=" * 80)

def convert_to_image(diagram_content, output_file, image_format, width, height):
    """Convert Mermaid diagram to image using mermaid-cli"""
    # Try different possible locations for mmdc
    mmdc_commands = ['mmdc', 'mmdc.cmd', r'C:\Users\cjony\AppData\Roaming\npm\mmdc', r'C:\Users\cjony\AppData\Roaming\npm\mmdc.cmd']
    
    mmdc_cmd = None
    for cmd in mmdc_commands:
        try:
            subprocess.run([cmd, '--version'], capture_output=True, check=True, timeout=5)
            mmdc_cmd = cmd
            break
        except (subprocess.CalledProcessError, FileNotFoundError, subprocess.TimeoutExpired):
            continue
    
    if not mmdc_cmd:
        click.echo("Error: mermaid-cli is not installed or not found in PATH.")
        click.echo("Please install it with: npm install -g @mermaid-js/mermaid-cli")
        click.echo("Or check if it's installed at: C:\\Users\\cjony\\AppData\\Roaming\\npm\\")
        return
    
    try:
        # Create temporary file for mermaid content
        with tempfile.NamedTemporaryFile(mode='w', suffix='.md', delete=False, encoding='utf-8', newline='') as temp_file:
            temp_file.write(diagram_content)
            temp_file_path = temp_file.name
        
        # Determine format from output file extension
        file_extension = output_file.split('.')[-1].lower()
        actual_format = file_extension if file_extension in ['png', 'svg', 'pdf'] else image_format
        
        # Convert to image using mermaid-cli
        cmd = [
            mmdc_cmd,
            '-i', temp_file_path,
            '-o', output_file,
            '-e', actual_format,
            '-w', str(width),
            '-H', str(height)
        ]
        
        # Add PDF-specific options
        if image_format.lower() == 'pdf':
            # Set background color for PDF
            cmd.extend(['-b', 'white'])
            # Add PDF fit option
            cmd.append('-f')
        
        click.echo(f"Running command: {' '.join(cmd)}")
        
        # Set environment variables for UTF-8 encoding
        env = os.environ.copy()
        env['PYTHONIOENCODING'] = 'utf-8'
        env['LANG'] = 'en_US.UTF-8'
        env['LC_ALL'] = 'en_US.UTF-8'
        
        result = subprocess.run(cmd, capture_output=True, text=True, check=True, encoding='utf-8', errors='ignore', env=env)
        
        # Clean up temporary file
        os.unlink(temp_file_path)
        
        # Check if the expected output file was created
        if os.path.exists(output_file):
            actual_format = output_file.split('.')[-1].upper()
            click.echo(f"Image saved to: {output_file}")
            click.echo(f"Format: {actual_format}, Size: {width}x{height}")
        else:
            # Check for files with similar names (mermaid-cli sometimes adds numbers)
            import glob
            pattern = output_file.replace('.pdf', '-*.pdf').replace('.png', '-*.png').replace('.svg', '-*.svg')
            matching_files = glob.glob(pattern)
            if matching_files:
                actual_file = matching_files[0]
                actual_format = actual_file.split('.')[-1].upper()
                click.echo(f"Image saved to: {actual_file}")
                click.echo(f"Format: {actual_format}, Size: {width}x{height}")
                click.echo(f"Note: mermaid-cli created {actual_file} instead of {output_file}")
            else:
                click.echo(f"Warning: Expected file {output_file} not found")
        
        click.echo(f"Command output: {result.stdout}")
        
    except subprocess.CalledProcessError as e:
        click.echo(f"Error converting to image: {e}")
        click.echo(f"Error output: {e.stderr}")
        if os.path.exists(temp_file_path):
            os.unlink(temp_file_path)
    except Exception as e:
        click.echo(f"Unexpected error: {e}")
        if 'temp_file_path' in locals() and os.path.exists(temp_file_path):
            os.unlink(temp_file_path)

def _save_crud_matrix_as_excel(matrix, project_name, output_path):
    """CRUD matrix 데이터를 Excel 파일로 저장"""
    try:
        import pandas as pd
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        
        # DataFrame 생성
        df_data = []
        for row in matrix:
            df_data.append({
                'Package': row['package_name'] or 'N/A',
                'Class Name': row['class_name'],
                'Method': row['method_name'],
                'Schema': row['schema'] or 'unknown',
                'Table': row['table_name'],
                'Operations': ', '.join(row['operations']) if row['operations'] else 'None'
            })
        
        df = pd.DataFrame(df_data)
        
        # Excel 파일 생성 및 스타일 적용
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='CRUD Matrix', index=False)
            
            # 워크시트 스타일 적용
            workbook = writer.book
            worksheet = writer.sheets['CRUD Matrix']
            
            # 헤더 스타일
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            
            # 컬럼 너비 자동 조정
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)
        
        return True
        
    except ImportError as e:
        click.echo(f"Error: Required library not found: {e}")
        click.echo("Please install required libraries: pip install pandas openpyxl")
        return False
    except Exception as e:
        click.echo(f"Error creating Excel file: {e}")
        return False

def _save_crud_matrix_as_image(matrix, project_name, output_path, image_format):
    """CRUD matrix 데이터를 SVG/PNG 이미지로 저장 (Enhanced classDiagram 형식)"""
    try:
        # Mermaid classDiagram 형식으로 변경
        diagram_lines = ["```mermaid", "classDiagram"]
        
        # 클래스별 데이터 구조화
        class_data = {}
        for row in matrix:
            class_name = row['class_name']
            if class_name not in class_data:
                class_data[class_name] = {
                    'package_name': row['package_name'],
                    'methods': []
                }
            
            # 메서드-테이블-CRUD 정보 저장
            schema = row['schema'] if row['schema'] and row['schema'] != 'unknown' else None
            table_display = f"{schema}.{row['table_name']}" if schema else row['table_name']
            operations = ', '.join(row['operations']) if row['operations'] else 'None'
            
            class_data[class_name]['methods'].append({
                'method_name': row['method_name'],
                'table': table_display,
                'operations': operations
            })
        
        # 각 클래스에 대해 classDiagram 노드 생성
        for class_name, data in class_data.items():
            # 패키지명 정리
            package_name = data['package_name'] or 'N/A'
            clean_class_name = class_name.replace('.', '_').replace('-', '_').replace(' ', '_')
            
            # 클래스 정의 시작
            diagram_lines.append(f"    class {clean_class_name} {{")
            diagram_lines.append(f"        <<{package_name}>>")
            
            # 메서드 목록 추가 (최대 10개)
            methods = data['methods'][:10]  # 최대 10개로 제한
            for method_info in methods:
                method_line = f"        +{method_info['method_name']}() {method_info['table']} [{method_info['operations']}]"
                diagram_lines.append(method_line)
            
            # 10개 이상이면 생략 표시
            if len(data['methods']) > 10:
                diagram_lines.append(f"        ... ({len(data['methods']) - 10} more)")
            
            diagram_lines.append("    }")
        
        # 테이블 노드 생성
        table_nodes = set()
        for row in matrix:
            schema = row['schema'] if row['schema'] and row['schema'] != 'unknown' else None
            table_display = f"{schema}.{row['table_name']}" if schema else row['table_name']
            table_key = table_display.replace('.', '_').replace('-', '_').replace(' ', '_')
            
            if table_key not in table_nodes:
                diagram_lines.append(f"    class {table_key} {{")
                diagram_lines.append(f"        <<Database Table>>")
                diagram_lines.append(f"        {table_display}")
                diagram_lines.append("    }")
                table_nodes.add(table_key)
        
        # 클래스-테이블 관계 생성
        for class_name, data in class_data.items():
            clean_class_name = class_name.replace('.', '_').replace('-', '_').replace(' ', '_')
            for method_info in data['methods'][:10]:  # 최대 10개 관계만 표시
                table_key = method_info['table'].replace('.', '_').replace('-', '_').replace(' ', '_')
                diagram_lines.append(f"    {clean_class_name} --> {table_key}")
        
        diagram_lines.append("```")
        diagram = "\n".join(diagram_lines)
        
        # 이미지로 변환 (크기 증가)
        convert_to_image(diagram, output_path, image_format, 2400, 1800)
        return True
        
    except Exception as e:
        click.echo(f"Error creating image file: {e}")
        return False

@click.group()
def cli():
    pass

def _validate_analyze_options(db_object, java_object, class_name, update, java_source_folder):
    """analyze 명령어 옵션 검증"""
    # Check if at least one analysis type is specified
    if not db_object and not java_object and not class_name and not update:
        # Default to full analysis (Java + DB) if no flags are specified
        java_object = True
        db_object = True
        click.echo("No analysis type specified. Defaulting to full analysis (Java + DB objects).")
        click.echo("Use --db-object, --java-object, --all-objects, --class-name, or --update to specify analysis type.")
    
    # Check java_source_folder requirement for Java-related operations
    if java_object or class_name or update or (not db_object):
        if not java_source_folder:
            click.echo("Error: JAVA_SOURCE_FOLDER environment variable or --java-source-folder option is required for Java object analysis.", err=True)
            exit(1)
    
    return db_object, java_object

def _determine_project_name(project_name, detected_name, logger):
    """프로젝트명 결정 (우선순위: --project-name > detected_name)"""
    if project_name:
        final_project_name = project_name
        logger.info(f"Using provided project name: {final_project_name}")
    else:
        final_project_name = detected_name
        logger.info(f"Using detected project name: {final_project_name}")
    
    return final_project_name

def _calculate_java_statistics(packages, classes, beans, endpoints, mappers, entities, repos, queries, config_files, test_classes, sql_statements):
    """Java 분석 통계 계산"""
    total_methods = sum(len(class_obj.methods) for class_obj in classes)
    total_fields = sum(len(class_obj.properties) for class_obj in classes)
    
    return {
        'packages': len(packages),
        'classes': len(classes),
        'methods': total_methods,
        'fields': total_fields,
        'beans': len(beans),
        'endpoints': len(endpoints),
        'mybatis_mappers': len(mappers),
        'jpa_entities': len(entities),
        'jpa_repositories': len(repos),
        'jpa_queries': len(queries),
        'config_files': len(config_files),
        'test_classes': len(test_classes),
        'sql_statements': len(sql_statements)
    }

def _print_analysis_summary(overall_start_time, java_stats, db_stats, logger):
    """분석 결과 요약 출력"""
    overall_end_time = datetime.now()
    overall_duration = (overall_end_time - overall_start_time).total_seconds()
    
    logger.info("")
    logger.info("=" * 80)
    logger.info("ANALYSIS SUMMARY")
    logger.info("=" * 80)
    
    if java_stats:
        logger.info(f"Java Analysis:")
        logger.info(f"  • Packages: {java_stats['packages']}")
        logger.info(f"  • Classes: {java_stats['classes']}")
        logger.info(f"  • Methods: {java_stats['methods']}")
        logger.info(f"  • Fields: {java_stats['fields']}")
        logger.info(f"  • Spring Beans: {java_stats['beans']}")
        logger.info(f"  • REST Endpoints: {java_stats['endpoints']}")
        logger.info(f"  • MyBatis Mappers: {java_stats['mybatis_mappers']}")
        logger.info(f"  • JPA Entities: {java_stats['jpa_entities']}")
        logger.info(f"  • JPA Repositories: {java_stats['jpa_repositories']}")
        logger.info(f"  • JPA Queries: {java_stats['jpa_queries']}")
        logger.info(f"  • Config Files: {java_stats['config_files']}")
        logger.info(f"  • Test Classes: {java_stats['test_classes']}")
        logger.info(f"  • SQL Statements: {java_stats['sql_statements']}")
    
    if db_stats:
        logger.info(f"Database Analysis:")
        logger.info(f"  • Databases: {db_stats['databases']}")
        logger.info(f"  • Tables: {db_stats['tables']}")
        logger.info(f"  • Columns: {db_stats['columns']}")
        logger.info(f"  • Indexes: {db_stats['indexes']}")
        logger.info(f"  • Constraints: {db_stats['constraints']}")
    
    logger.info(f"Total Analysis Time: {format_duration(overall_duration)}")
    logger.info("=" * 80)
    logger.info("====== analyze 작업 완료 ======")
    logger.info("")

def _analyze_full_project_java(java_source_folder, project_name, logger):
    """Full Project Analysis - Java 파싱 및 프로젝트명 결정"""
    click.echo(f"Parsing Java project at: {java_source_folder}")
    packages_to_add, classes_to_add, class_to_package_map, beans, dependencies, endpoints, mybatis_mappers, jpa_entities, jpa_repositories, jpa_queries, config_files, test_classes, sql_statements, detected_project_name = parse_java_project(java_source_folder)
    
    # 프로젝트명 결정
    final_project_name = _determine_project_name(project_name, detected_project_name, logger)
    
    logger.info(f"Project name: {final_project_name}")
    logger.info(f"Found {len(packages_to_add)} packages and {len(classes_to_add)} classes.")
    
    return packages_to_add, classes_to_add, class_to_package_map, beans, dependencies, endpoints, mybatis_mappers, jpa_entities, jpa_repositories, jpa_queries, config_files, test_classes, sql_statements, final_project_name

def _save_java_objects_to_neo4j(db, packages_to_add, classes_to_add, class_to_package_map, beans, dependencies, endpoints, mybatis_mappers, jpa_entities, jpa_repositories, jpa_queries, config_files, test_classes, sql_statements, final_project_name, clean, concurrent, workers, logger):
    """Full Project Analysis - Java 객체들을 Neo4j에 저장"""
    java_start_time = datetime.now()
    
    # Java Object 분석 통계 계산
    java_stats_temp = _calculate_java_statistics(packages_to_add, classes_to_add, beans, endpoints, mybatis_mappers, jpa_entities, jpa_repositories, jpa_queries, config_files, test_classes, sql_statements)
    java_stats_temp['project_name'] = final_project_name
    
    # Connect to database if db is None
    if db is None:
        logger.info(f"Connecting to Neo4j...")
        neo4j_uri = os.getenv("NEO4J_URI", "bolt://localhost:7687")
        neo4j_user = os.getenv("NEO4J_USER", "neo4j")
        neo4j_password = os.getenv("NEO4J_PASSWORD")
        neo4j_database = os.getenv("NEO4J_DATABASE", "neo4j")
        
        if not neo4j_password:
            logger.error("NEO4J_PASSWORD not set - cannot connect to database")
            raise ValueError("NEO4J_PASSWORD environment variable is required")
        
        db, pool = _connect_to_neo4j_db(neo4j_uri, neo4j_user, neo4j_password, neo4j_database, logger)
    
    # Add project
    logger.info(f"DB 저장 -  project: {final_project_name}")
    project = Project(
        name=final_project_name,
        display_name=final_project_name,
        description=f"Java project: {final_project_name}",
        repository_url="",
        language="Java",
        framework="Spring Boot",
        version="1.0",
        ai_description="",
        created_at=datetime.now().strftime("%Y/%m/%d %H:%M:%S.%f")[:-3]
    )
    db.add_project(project)
    
    # Add packages
    logger.info(f"DB 저장 -  {len(packages_to_add)} packages...")
    for package in packages_to_add:
        db.add_package(package, final_project_name)
    
    # Add classes with concurrent processing if enabled
    if concurrent:
        logger.info(f"DB 저장 -  {len(classes_to_add)} classes with concurrent processing...")
        db.add_classes_concurrent(classes_to_add, class_to_package_map, final_project_name, workers)
    else:
        total = len(classes_to_add)
        logger.info(f"DB 저장 -  {total} classes...")
        last_percent = 0
        for idx, class_obj in enumerate(classes_to_add, 1):
            package_name = class_to_package_map.get(class_obj.name, "unknown")
            db.add_class(class_obj, package_name, final_project_name)
            
            percent = int((idx / total) * 100)
            if percent >= last_percent + 10 or idx == total:
                last_percent = percent
                logger.info(f"   - classes 저장중 [{idx}/{total}] ({percent}%)")
    
    # Add Spring Beans
    logger.info(f"DB 저장 -  {len(beans)} Spring Beans...")
    for bean in beans:
        db.add_bean(bean, final_project_name)
    
    # Add Bean Dependencies
    logger.info(f"DB 저장 -  {len(dependencies)} Bean Dependencies...")
    for dep in dependencies:
        db.add_bean_dependency(dep, final_project_name)
    
    # Add REST Endpoints
    logger.info(f"DB 저장 -  {len(endpoints)} REST Endpoints...")
    for endpoint in endpoints:
        db.add_endpoint(endpoint, final_project_name)
    
    # Add MyBatis Mappers
    logger.info(f"DB 저장 -  {len(mybatis_mappers)} MyBatis Mappers...")
    for mapper in mybatis_mappers:
        db.add_mybatis_mapper(mapper, final_project_name)
    
    # Add JPA Entities
    logger.info(f"DB 저장 -  {len(jpa_entities)} JPA Entities...")
    for entity in jpa_entities:
        db.add_jpa_entity(entity, final_project_name)
    
    # Add JPA Repositories
    logger.info(f"DB 저장 -  {len(jpa_repositories)} JPA Repositories...")
    for repo in jpa_repositories:
        db.add_jpa_repository(repo, final_project_name)
    
    # Add JPA Queries
    logger.info(f"DB 저장 -  {len(jpa_queries)} JPA Queries...")
    for query in jpa_queries:
        db.add_jpa_query(query, final_project_name)
    
    # Add Config Files
    logger.info(f"DB 저장 -  {len(config_files)} Config Files...")
    for config in config_files:
        db.add_config_file(config, final_project_name)
    
    # Add Test Classes
    logger.info(f"DB 저장 -  {len(test_classes)} Test Classes...")
    for test_class in test_classes:
        db.add_test_class(test_class, final_project_name)
    
    # Add SQL Statements
    logger.info(f"DB 저장 -  {len(sql_statements)} SQL Statements...")
    for sql_stmt in sql_statements:
        db.add_sql_statement(sql_stmt, final_project_name)
    
    # Java Object 분석 완료 시각 기록
    java_end_time = datetime.now()
    java_duration = (java_end_time - java_start_time).total_seconds()
    
    logger.info(f"Java object analysis completed in {format_duration(java_duration)}")
    
    # 시간 정보 추가
    java_stats_temp['start_time'] = java_start_time
    java_stats_temp['end_time'] = java_end_time
    
    return java_stats_temp

def _analyze_full_project_db(db, db_script_folder, final_project_name, dry_run, logger):
    """Full Project Analysis - DB 객체 분석 및 저장"""
    if not db_script_folder:
        logger.info("No DB_SCRIPT_FOLDER specified. Skipping database object analysis.")
        return None
    
    logger.info(f"Analyzing database objects from: {db_script_folder}")
    
    # Parse DDL files
    db_parser = DBParser()
    all_db_objects = db_parser.parse_ddl_directory(db_script_folder, final_project_name)
    
    if not all_db_objects:
        logger.info("No database objects found.")
        return None
    
    logger.info(f"Found {len(all_db_objects)} DDL files to process.")
    
    # Group objects by database
    grouped_objects = {}
    for obj in all_db_objects:
        db_name = obj['database'].name or "default"
        if db_name not in grouped_objects:
            grouped_objects[db_name] = []
        grouped_objects[db_name].append(obj)
    
    logger.info(f"Found {len(grouped_objects)} databases: {list(grouped_objects.keys())}")
    
    if dry_run:
        click.echo("Dry run mode - not saving to database.")
        for db_name, objects in grouped_objects.items():
            click.echo(f"Database '{db_name}': {len(objects)} objects")
        click.echo("Database object analysis complete (dry run).")
        return None
    
    # Save to Neo4j
    db_stats = {
        'databases': len(grouped_objects),
        'tables': 0,
        'columns': 0,
        'indexes': 0,
        'constraints': 0
    }
    
    for db_name, objects in grouped_objects.items():
        logger.info(f"Processing database: {db_name}")
        
        for obj in objects:
            # Add database
            db.add_database(obj['database'], final_project_name)
            
            # Add tables
            for table in obj['tables']:
                db.add_table(table, db_name, final_project_name)
                db_stats['tables'] += 1
            
            # Add columns
            for column in obj['columns']:
                db.add_column(column, column.table_name, final_project_name)
                db_stats['columns'] += 1
            
            # Add indexes
            for index, table_name in obj['indexes']:
                db.add_index(index, table_name, final_project_name)
                db_stats['indexes'] += 1
            
            # Add constraints
            for constraint, table_name in obj['constraints']:
                db.add_constraint(constraint, table_name, final_project_name)
                db_stats['constraints'] += 1
    
    logger.info(f"Database object analysis completed.")
    return db_stats

# =============================================================================
# 공통 유틸리티 함수들 (모듈화를 위한 헬퍼 함수들)
# =============================================================================

def _parse_java_with_concurrency(java_source_folder, concurrent, workers, logger):
    """Java 소스 파싱 (concurrent 옵션 처리 포함)"""
    if concurrent:
        # TODO: parse_java_project_concurrent 함수가 구현되면 다시 활성화
        logger.warning("Concurrent processing requested but not yet implemented. Using single-threaded processing.")
        logger.info("Using single-threaded processing")
        return parse_java_project(java_source_folder)
    else:
        return parse_java_project(java_source_folder)

def _get_or_determine_project_name(project_name, detected_project_name, java_source_folder, logger):
    """프로젝트명 결정 로직 통합"""
    if project_name:
        final_project_name = project_name
        logger.info(f"Using provided project name: {final_project_name}")
    elif detected_project_name:
        final_project_name = detected_project_name
        logger.info(f"Using detected project name: {final_project_name}")
    else:
        # Fallback: use folder name or default
        from pathlib import Path
        if java_source_folder:
            final_project_name = Path(java_source_folder).resolve().name
        else:
            final_project_name = os.getenv("PROJECT_NAME", "default_project")
        logger.info(f"Using fallback project name: {final_project_name}")
    
    return final_project_name

def _connect_to_neo4j_db(neo4j_uri, neo4j_user, neo4j_password, neo4j_database, logger):
    """Neo4j 연결 및 초기화"""
    logger.info(f"Connecting to Neo4j at {neo4j_uri}...")
    
    # Connection Pool 초기화
    pool = get_connection_pool()
    if not pool.is_initialized():
        pool_size = int(os.getenv('NEO4J_POOL_SIZE', '10'))
        logger.info(f"Initializing Neo4j connection pool with {pool_size} connections...")
        pool.initialize(neo4j_uri, neo4j_user, neo4j_password, neo4j_database, pool_size)
        logger.info(f"Connected to Neo4j at {neo4j_uri} (database: {neo4j_database})")
    else:
        logger.info(f"Using existing connection pool (database: {neo4j_database})")
    
    db = GraphDB(neo4j_uri, neo4j_user, neo4j_password, neo4j_database)
    return db, pool

def _clean_java_objects(db, logger):
    """Java 객체 정리"""
    logger.info("Cleaning Java objects...")
    conn = db._pool.acquire() if hasattr(db, '_pool') else None
    try:
        if conn:
            with conn.session() as session:
                session.run("MATCH (n:Package) DETACH DELETE n")
                session.run("MATCH (n:Class) DETACH DELETE n")
                session.run("MATCH (n:Method) DETACH DELETE n")
                session.run("MATCH (n:Field) DETACH DELETE n")
                session.run("MATCH (n:Bean) DETACH DELETE n")
                session.run("MATCH (n:Endpoint) DETACH DELETE n")
                session.run("MATCH (n:MyBatisMapper) DETACH DELETE n")
                session.run("MATCH (n:JpaEntity) DETACH DELETE n")
                session.run("MATCH (n:ConfigFile) DETACH DELETE n")
                session.run("MATCH (n:TestClass) DETACH DELETE n")
                session.run("MATCH (n:SqlStatement) DETACH DELETE n")
        else:
            with db._driver.session() as session:
                session.run("MATCH (n:Package) DETACH DELETE n")
                session.run("MATCH (n:Class) DETACH DELETE n")
                session.run("MATCH (n:Method) DETACH DELETE n")
                session.run("MATCH (n:Field) DETACH DELETE n")
                session.run("MATCH (n:Bean) DETACH DELETE n")
                session.run("MATCH (n:Endpoint) DETACH DELETE n")
                session.run("MATCH (n:MyBatisMapper) DETACH DELETE n")
                session.run("MATCH (n:JpaEntity) DETACH DELETE n")
                session.run("MATCH (n:ConfigFile) DETACH DELETE n")
                session.run("MATCH (n:TestClass) DETACH DELETE n")
                session.run("MATCH (n:SqlStatement) DETACH DELETE n")
    finally:
        if conn and hasattr(db, '_pool'):
            db._pool.release(conn)

def _clean_db_objects(db, logger):
    """DB 객체 정리"""
    logger.info("Cleaning database objects...")
    conn = db._pool.acquire() if hasattr(db, '_pool') else None
    try:
        if conn:
            with conn.session() as session:
                session.run("MATCH (n:Database) DETACH DELETE n")
                session.run("MATCH (n:Table) DETACH DELETE n")
                session.run("MATCH (n:Column) DETACH DELETE n")
                session.run("MATCH (n:Index) DETACH DELETE n")
                session.run("MATCH (n:Constraint) DETACH DELETE n")
        else:
            with db._driver.session() as session:
                session.run("MATCH (n:Database) DETACH DELETE n")
                session.run("MATCH (n:Table) DETACH DELETE n")
                session.run("MATCH (n:Column) DETACH DELETE n")
                session.run("MATCH (n:Index) DETACH DELETE n")
                session.run("MATCH (n:Constraint) DETACH DELETE n")
    finally:
        if conn and hasattr(db, '_pool'):
            db._pool.release(conn)

def _add_springboot_objects(db, beans, dependencies, endpoints, mybatis_mappers, jpa_entities, jpa_repositories, jpa_queries, config_files, test_classes, sql_statements, final_project_name, logger):
    """Spring Boot 관련 객체 저장"""
    import time
    
    # Add Spring Beans
    if beans:
        logger.info(f"DB 저장 -  {len(beans)} Spring Beans to database...")
        start_time = time.time()
        last_percent = 0
        for i, bean in enumerate(beans):
            db.add_bean(bean, final_project_name)
            percent = int(((i + 1) / len(beans)) * 100)
            if percent >= last_percent + 10 or (i + 1) == len(beans):
                last_percent = percent
                logger.info(f"   - beans 저장중 [{i+1}/{len(beans)}] ({percent}%)")
        logger.info(f"✓ Added {len(beans)} Spring Beans in {time.time() - start_time:.2f}s")
    
    # Add Bean Dependencies
    if dependencies:
        logger.info(f"DB 저장 -  {len(dependencies)} Bean dependencies to database...")
        start_time = time.time()
        for dependency in dependencies:
            db.add_bean_dependency(dependency, final_project_name)
        logger.info(f"✓ Added {len(dependencies)} Bean dependencies in {time.time() - start_time:.2f}s")
    
    # Add REST Endpoints
    if endpoints:
        logger.info(f"DB 저장 -  {len(endpoints)} REST API endpoints to database...")
        start_time = time.time()
        last_percent = 0
        for i, endpoint in enumerate(endpoints):
            db.add_endpoint(endpoint, final_project_name)
            percent = int(((i + 1) / len(endpoints)) * 100)
            if percent >= last_percent + 10 or (i + 1) == len(endpoints):
                last_percent = percent
                logger.info(f"   - endpoints 저장중 [{i+1}/{len(endpoints)}] ({percent}%)")
        logger.info(f"✓ Added {len(endpoints)} REST API endpoints in {time.time() - start_time:.2f}s")
    
    # Add MyBatis Mappers
    if mybatis_mappers:
        logger.info(f"DB 저장 -  {len(mybatis_mappers)} MyBatis mappers to database...")
        start_time = time.time()
        for mapper in mybatis_mappers:
            db.add_mybatis_mapper(mapper, final_project_name)
        logger.info(f"✓ Added {len(mybatis_mappers)} MyBatis mappers in {time.time() - start_time:.2f}s")
    
    # Add JPA Entities
    if jpa_entities:
        logger.info(f"DB 저장 -  {len(jpa_entities)} JPA entities to database...")
        start_time = time.time()
        for entity in jpa_entities:
            db.add_jpa_entity(entity, final_project_name)
        logger.info(f"✓ Added {len(jpa_entities)} JPA entities in {time.time() - start_time:.2f}s")
    
    # Add JPA Repositories
    if jpa_repositories:
        logger.info(f"DB 저장 -  {len(jpa_repositories)} JPA repositories to database...")
        start_time = time.time()
        for repository in jpa_repositories:
            db.add_jpa_repository(repository, final_project_name)
        logger.info(f"✓ Added {len(jpa_repositories)} JPA repositories in {time.time() - start_time:.2f}s")
    
    # Add JPA Queries
    if jpa_queries:
        logger.info(f"DB 저장 -  {len(jpa_queries)} JPA queries to database...")
        start_time = time.time()
        last_percent = 0
        for i, query in enumerate(jpa_queries):
            db.add_jpa_query(query, final_project_name)
            percent = int(((i + 1) / len(jpa_queries)) * 100)
            if percent >= last_percent + 10 or (i + 1) == len(jpa_queries):
                last_percent = percent
                logger.info(f"   - jpa_queries 저장중 [{i+1}/{len(jpa_queries)}] ({percent}%)")
        logger.info(f"✓ Added {len(jpa_queries)} JPA queries in {time.time() - start_time:.2f}s")
    
    # Add Config Files
    if config_files:
        logger.info(f"DB 저장 -  {len(config_files)} configuration files to database...")
        start_time = time.time()
        for config_file in config_files:
            db.add_config_file(config_file, final_project_name)
        logger.info(f"✓ Added {len(config_files)} configuration files in {time.time() - start_time:.2f}s")
    
    # Add Test Classes
    if test_classes:
        logger.info(f"DB 저장 -  {len(test_classes)} test classes to database...")
        start_time = time.time()
        for test_class in test_classes:
            db.add_test_class(test_class, final_project_name)
        logger.info(f"✓ Added {len(test_classes)} test classes in {time.time() - start_time:.2f}s")
    
    # Add SQL Statements
    if sql_statements:
        logger.info(f"DB 저장 -  {len(sql_statements)} SQL statements to database...")
        start_time = time.time()
        last_percent = 0
        for i, sql_statement in enumerate(sql_statements):
            db.add_sql_statement(sql_statement, final_project_name)
            # Create relationship between mapper and SQL statement
            conn = db._pool.acquire() if hasattr(db, '_pool') else None
            try:
                if conn:
                    with conn.session() as session:
                        session.execute_write(db._create_mapper_sql_relationship_tx, sql_statement.mapper_name, sql_statement.id, final_project_name)
                else:
                    with db._driver.session() as session:
                        session.execute_write(db._create_mapper_sql_relationship_tx, sql_statement.mapper_name, sql_statement.id, final_project_name)
            finally:
                if conn and hasattr(db, '_pool'):
                    db._pool.release(conn)
            
            percent = int(((i + 1) / len(sql_statements)) * 100)
            if percent >= last_percent + 10 or (i + 1) == len(sql_statements):
                last_percent = percent
                logger.info(f"   - sql_statements 저장중 [{i+1}/{len(sql_statements)}] ({percent}%)")
        logger.info(f"✓ Added {len(sql_statements)} SQL statements in {time.time() - start_time:.2f}s")

def _add_single_class_objects(db, class_node, package_name, final_project_name, logger):
    """단일 클래스의 모든 관련 객체 저장"""
    # Extract and add related Spring Boot analysis results for this class only
    classes_list = [class_node]
    beans = extract_beans_from_classes(classes_list)
    dependencies = analyze_bean_dependencies(classes_list, beans)
    endpoints = extract_endpoints_from_classes(classes_list)
    mybatis_mappers = extract_mybatis_mappers_from_classes(classes_list)
    jpa_entities = extract_jpa_entities_from_classes(classes_list)
    test_classes = extract_test_classes_from_classes(classes_list)
    
    # Extract SQL statements from MyBatis mappers
    sql_statements = extract_sql_statements_from_mappers(mybatis_mappers, final_project_name)
    
    # Add Spring Boot analysis results
    if beans:
        logger.info(f"DB 저장 -  {len(beans)} Spring Beans to database...")
        for bean in beans:
            db.add_bean(bean, final_project_name)
    
    if dependencies:
        logger.info(f"DB 저장 -  {len(dependencies)} Bean dependencies to database...")
        for dependency in dependencies:
            db.add_bean_dependency(dependency, final_project_name)
    
    if endpoints:
        logger.info(f"DB 저장 -  {len(endpoints)} REST API endpoints to database...")
        for endpoint in endpoints:
            db.add_endpoint(endpoint, final_project_name)
    
    if mybatis_mappers:
        logger.info(f"DB 저장 -  {len(mybatis_mappers)} MyBatis mappers to database...")
        for mapper in mybatis_mappers:
            db.add_mybatis_mapper(mapper, final_project_name)
    
    if jpa_entities:
        logger.info(f"DB 저장 -  {len(jpa_entities)} JPA entities to database...")
        for entity in jpa_entities:
            db.add_jpa_entity(entity, final_project_name)
    
    if test_classes:
        logger.info(f"DB 저장 -  {len(test_classes)} test classes to database...")
        for test_class in test_classes:
            db.add_test_class(test_class, final_project_name)
    
    if sql_statements:
        logger.info(f"DB 저장 -  {len(sql_statements)} SQL statements to database...")
        for sql_statement in sql_statements:
            db.add_sql_statement(sql_statement, final_project_name)
            # Create relationship between mapper and SQL statement
            conn = db._pool.acquire() if hasattr(db, '_pool') else None
            try:
                if conn:
                    with conn.session() as session:
                        session.execute_write(db._create_mapper_sql_relationship_tx, sql_statement.mapper_name, sql_statement.id, final_project_name)
                else:
                    with db._driver.session() as session:
                        session.execute_write(db._create_mapper_sql_relationship_tx, sql_statement.mapper_name, sql_statement.id, final_project_name)
            finally:
                if conn and hasattr(db, '_pool'):
                    db._pool.release(conn)

# =============================================================================
# 시나리오별 핸들러 함수들 (각 분석 시나리오를 담당하는 함수들)
# =============================================================================

def _handle_full_project_analysis(java_source_folder, project_name, neo4j_uri, neo4j_user, neo4j_password, neo4j_database, clean, dry_run, concurrent, workers, logger):
    """전체 프로젝트 분석 (Java + DB) 핸들러"""
    # Java 분석
    packages_to_add, classes_to_add, class_to_package_map, beans, dependencies, endpoints, mybatis_mappers, jpa_entities, jpa_repositories, jpa_queries, config_files, test_classes, sql_statements, final_project_name = _analyze_full_project_java(java_source_folder, project_name, logger)
    
    # Java 객체들을 Neo4j에 저장
    java_stats = _save_java_objects_to_neo4j(None, packages_to_add, classes_to_add, class_to_package_map, beans, dependencies, endpoints, mybatis_mappers, jpa_entities, jpa_repositories, jpa_queries, config_files, test_classes, sql_statements, final_project_name, clean, concurrent, workers, logger)
    
    # DB 객체 분석
    db_script_folder = os.getenv("DB_SCRIPT_FOLDER")
    db_stats = _analyze_full_project_db(None, db_script_folder, final_project_name, dry_run, logger)
    
    return java_stats, db_stats

def _handle_java_only_analysis(java_source_folder, project_name, neo4j_uri, neo4j_user, neo4j_password, neo4j_database, clean, dry_run, concurrent, workers, logger):
    """Java 객체만 분석 핸들러"""
    logger.info("Analyzing Java objects from source code...")
    
    # Java Object 분석 시작 시각 기록
    java_start_time = datetime.now()
    
    if not java_source_folder:
        logger.error("Error: JAVA_SOURCE_FOLDER environment variable is required for --java_object option.", err=True)
        logger.error("Please set JAVA_SOURCE_FOLDER in your .env file or environment variables.")
        exit(1)
    
    if not os.path.exists(java_source_folder):
        logger.error(f"Error: Java source folder {java_source_folder} does not exist.", err=True)
        exit(1)
    
    try:
        # Parse Java project
        logger.info(f"Parsing Java project at: {java_source_folder}")
        packages_to_add, classes_to_add, class_to_package_map, beans, dependencies, endpoints, mybatis_mappers, jpa_entities, jpa_repositories, jpa_queries, config_files, test_classes, sql_statements, detected_project_name = _parse_java_with_concurrency(java_source_folder, concurrent, workers, logger)
        
        # 프로젝트명 결정
        final_project_name = _get_or_determine_project_name(project_name, detected_project_name, java_source_folder, logger)
        
        logger.info(f"Project name: {final_project_name}")
        logger.info(f"Found {len(packages_to_add)} packages and {len(classes_to_add)} classes.")
        
        # Java Object 분석 완료 시각 기록 및 통계 수집
        java_end_time = datetime.now()
        
        # Methods와 Fields 개수 계산
        total_methods = sum(len(class_obj.methods) for class_obj in classes_to_add)
        total_fields = sum(len(class_obj.properties) for class_obj in classes_to_add)
        
        # Java Object 분석 통계는 DB 저장 완료 후에 수집하므로 여기서는 임시 저장
        java_stats_temp = {
            'project_name': final_project_name,
            'total_files': 0,  # parse 함수에서 반환하지 않으므로 0으로 설정
            'processed_files': len(classes_to_add),  # 클래스 개수로 대체
            'error_files': 0,  # parse 함수에서 반환하지 않으므로 0으로 설정
            'packages': len(packages_to_add),
            'classes': len(classes_to_add),
            'methods': total_methods,
            'fields': total_fields,
            'beans': len(beans),
            'endpoints': len(endpoints),
            'mybatis_mappers': len(mybatis_mappers),
            'jpa_entities': len(jpa_entities),
            'jpa_repositories': len(jpa_repositories),
            'sql_statements': len(sql_statements),
        }
        
        if dry_run:
            click.echo("Dry run mode - not connecting to database.")
            click.echo(f"Found {len(packages_to_add)} packages and {len(classes_to_add)} classes.")
            click.echo(f"Found {len(beans)} Spring Beans and {len(dependencies)} dependencies.")
            click.echo(f"Found {len(endpoints)} REST API endpoints.")
            click.echo(f"Found {len(mybatis_mappers)} MyBatis mappers.")
            click.echo(f"Found {len(jpa_entities)} JPA entities.")
            click.echo(f"Found {len(jpa_repositories)} JPA repositories.")
            click.echo(f"Found {len(jpa_queries)} JPA queries.")
            click.echo(f"Found {len(config_files)} configuration files.")
            click.echo(f"Found {len(test_classes)} test classes.")
            click.echo(f"Found {len(sql_statements)} SQL statements.")
            click.echo("Java object analysis complete (dry run).")
            
            # Java Object 분석 통계 완성 (dry-run 모드)
            java_stats = {
                'start_time': java_start_time,
                'end_time': java_end_time,
                **java_stats_temp
            }
            
            return java_stats, None
        
        # Connect to database
        db, pool = _connect_to_neo4j_db(neo4j_uri, neo4j_user, neo4j_password, neo4j_database, logger)
        
        # Create or update Project node
        from csa.models.graph_entities import Project
        project_node = Project(
            name=final_project_name,
            display_name=final_project_name,
            language="Java",
        )
        db.add_project(project_node)
        logger.info(f"Project node created/updated: {final_project_name}")
        
        if clean:
            _clean_java_objects(db, logger)
        
        # Add packages
        logger.info("Adding packages to database...")
        for package_node in packages_to_add:
            db.add_package(package_node, final_project_name)
        
        # Add classes
        logger.info("Adding classes to database...")
        logger.info(f"Total classes to add: {len(classes_to_add)}")
        
        import time
        start_time = time.time()
        last_percent = 0
        
        for i, class_node in enumerate(classes_to_add):
            try:
                # Find the package for this class using the mapping
                class_key = f"{class_node.package_name}.{class_node.name}"
                package_name = class_to_package_map.get(class_key, class_node.package_name)
                
                if not package_name:
                    package_name = class_node.package_name
                
                class_start_time = time.time()
                logger.debug(f"Adding class {i+1}/{len(classes_to_add)}: {class_node.name} (package: {package_name})")
                db.add_class(class_node, package_name, final_project_name)
                
                # 메모리 절약: concurrent 옵션 사용 시 source 필드 제거
                if concurrent:
                    # TODO: clear_class_source_from_memory 함수가 구현되면 다시 활성화
                    # 현재는 간단히 source 필드를 None으로 설정
                    class_node.source = None
                    logger.debug(f"Cleared source from memory for class: {class_node.name}")
                
                class_elapsed = time.time() - class_start_time
                if class_elapsed > 1.0:  # 1초 이상 걸린 경우에만 시간 표시
                    logger.debug(f"  ✓ Completed in {class_elapsed:.2f}s")
                
                # 10% 단위로 전체 진행상태 표시
                percent = int(((i + 1) / len(classes_to_add)) * 100)
                if percent >= last_percent + 10 or (i + 1) == len(classes_to_add):
                    last_percent = percent
                    elapsed = time.time() - start_time
                    remaining = (elapsed / (i + 1)) * (len(classes_to_add) - i - 1) if (i + 1) < len(classes_to_add) else 0
                    logger.info(f"   - classes 저장중 [{i+1}/{len(classes_to_add)}] ({percent}%) - {elapsed:.1f}s elapsed, ~{remaining:.1f}s remaining")
                
            except Exception as e:
                click.echo(f"Error adding class {class_node.name}: {e}")
                continue
        
        total_elapsed = time.time() - start_time
        logger.info(f"✓ All {len(classes_to_add)} classes added successfully in {total_elapsed:.2f}s")
        
        # Add Spring Boot analysis results
        _add_springboot_objects(db, beans, dependencies, endpoints, mybatis_mappers, jpa_entities, jpa_repositories, jpa_queries, config_files, test_classes, sql_statements, final_project_name, logger)
        
        db.close()
        logger.info("Java object analysis complete.")
        
        # Java Object 분석 통계 완성
        java_stats = {
            'start_time': java_start_time,
            'end_time': java_end_time,
            **java_stats_temp
        }
        
        return java_stats, None
        
    except Exception as e:
        import traceback
        click.echo(f"Error analyzing Java objects: {e}")
        click.echo(f"\nFull traceback:")
        traceback.print_exc()
        click.echo("Use --dry-run flag to parse without database connection.")
        exit(1)

def _handle_db_only_analysis(project_name, neo4j_uri, neo4j_user, neo4j_password, neo4j_database, dry_run, logger):
    """DB 객체만 분석 핸들러"""
    logger.info("Analyzing database objects from DDL scripts...")
    
    # DB Object 분석 시작 시각 기록
    db_start_time = datetime.now()
    
    # Determine project name for DB analysis
    final_project_name = _get_or_determine_project_name(project_name, None, None, logger)
    
    db_script_folder = os.getenv("DB_SCRIPT_FOLDER")
    if not db_script_folder:
        logger.error("Error: DB_SCRIPT_FOLDER environment variable is required for --db_object option.", err=True)
        logger.error("Please set DB_SCRIPT_FOLDER in your .env file or environment variables.")
        exit(1)
    
    if not os.path.exists(db_script_folder):
        logger.error(f"Error: DB script folder {db_script_folder} does not exist.", err=True)
        exit(1)
    
    try:
        # Connect to database
        db, pool = _connect_to_neo4j_db(neo4j_uri, neo4j_user, neo4j_password, neo4j_database, logger)
        
        # DB 객체 분석
        db_stats = _analyze_full_project_db(db, db_script_folder, final_project_name, dry_run, logger)
        
        db.close()
        logger.info("Database object analysis complete.")
        
        return None, db_stats
        
    except Exception as e:
        logger.error(f"Error analyzing database objects: {e}")
        logger.error("Use --dry-run flag to parse without database connection.")
        exit(1)

def _handle_specific_class_analysis(java_source_folder, class_name, project_name, neo4j_uri, neo4j_user, neo4j_password, neo4j_database, dry_run, logger):
    """특정 클래스 분석 핸들러"""
    logger.info(f"Analyzing specific class: {class_name}")
    
    # Determine project name for class analysis
    final_project_name = _get_or_determine_project_name(project_name, None, java_source_folder, logger)
    
    # Find the Java file for this class
    java_file_path = None
    for root, _, files in os.walk(java_source_folder):
        for file in files:
            if file.endswith(".java") and file.replace(".java", "") == class_name:
                java_file_path = os.path.join(root, file)
                break
        if java_file_path:
            break
    
    if not java_file_path:
        logger.error(f"Error: Could not find Java file for class '{class_name}'", err=True)
        exit(1)
    
    logger.info(f"Found Java file: {java_file_path}")
    
    try:
        # Parse the single Java file
        from csa.services.java_parser import parse_single_java_file
        
        package_node, class_node, package_name = parse_single_java_file(java_file_path, final_project_name)
        
        if package_node is None or class_node is None:
            logger.error(f"Error: Failed to parse Java file: {java_file_path}", err=True)
            logger.error("Please check if the file contains valid Java code.")
            exit(1)
        
        logger.info(f"Parsed class: {class_node.name}")
        logger.info(f"Package: {package_name}")
        logger.info(f"Methods: {len(class_node.methods)}")
        logger.info(f"Properties: {len(class_node.properties)}")
        logger.info(f"Method calls: {len(class_node.calls)}")
        
        if dry_run:
            logger.info("Dry run mode - not connecting to database.")
            logger.info("Analysis complete (dry run).")
            logger.info("====== analyze 작업 완료 ======")
            return None, None
        
        # Connect to database
        db, pool = _connect_to_neo4j_db(neo4j_uri, neo4j_user, neo4j_password, neo4j_database, logger)
        
        # Delete existing data for this class
        click.echo(f"Deleting existing data for class '{class_name}'...")
        db.delete_class_and_related_data(class_name, final_project_name)
        
        # Add package
        logger.info("Adding package to database...")
        db.add_package(package_node, final_project_name)
        
        # Add class
        logger.info("Adding class to database...")
        db.add_class(class_node, package_name, final_project_name)
        
        # Add related objects
        _add_single_class_objects(db, class_node, package_name, final_project_name, logger)
        
        db.close()
        logger.info("Class analysis complete.")
        logger.info("====== analyze 작업 완료 ======")
        
        return None, None
        
    except Exception as e:
        logger.error(f"Error analyzing class: {e}")
        logger.error("Use --dry-run flag to parse without database connection.")
        exit(1)

def _handle_update_classes(java_source_folder, project_name, neo4j_uri, neo4j_user, neo4j_password, neo4j_database, dry_run, logger):
    """클래스 업데이트 핸들러"""
    logger.info("Updating all classes individually...")
    
    # Determine project name for update analysis
    final_project_name = _get_or_determine_project_name(project_name, None, java_source_folder, logger)
    
    # Find all Java files
    java_files = []
    for root, _, files in os.walk(java_source_folder):
        for file in files:
            if file.endswith(".java"):
                java_files.append(os.path.join(root, file))
    
    if not java_files:
        logger.error("No Java files found in the specified directory.", err=True)
        exit(1)
    
    logger.info(f"Found {len(java_files)} Java files to process.")
    
    if dry_run:
        logger.info("Dry run mode - not connecting to database.")
        for java_file in java_files:
            try:
                from csa.services.java_parser import parse_single_java_file
                package_node, class_node, package_name = parse_single_java_file(java_file, final_project_name)
                logger.info(f"  {class_node.name} ({package_name}) - Methods: {len(class_node.methods)}, Properties: {len(class_node.properties)}")
            except Exception as e:
                logger.error(f"  Error parsing {java_file}: {e}")
        logger.info("Update analysis complete (dry run).")
        logger.info("====== analyze 작업 완료 ======")
        return None, None
    
    try:
        # Connect to database
        db, pool = _connect_to_neo4j_db(neo4j_uri, neo4j_user, neo4j_password, neo4j_database, logger)
        
        processed_count = 0
        error_count = 0
        
        for java_file in java_files:
            try:
                logger.debug(f"Processing: {java_file}")
                
                # Parse the single Java file
                from csa.services.java_parser import parse_single_java_file
                
                package_node, class_node, package_name = parse_single_java_file(java_file, final_project_name)
                
                logger.debug(f"  Parsed class: {class_node.name} (Package: {package_name})")
                
                # Delete existing data for this class
                logger.info(f"  Deleting existing data for class '{class_node.name}'...")
                db.delete_class_and_related_data(class_node.name, final_project_name)
                
                # Add package
                db.add_package(package_node, final_project_name)
                
                # Add class
                db.add_class(class_node, package_name, final_project_name)
                
                # Add related objects
                _add_single_class_objects(db, class_node, package_name, final_project_name, logger)
                
                processed_count += 1
                logger.info(f"  [OK] Successfully processed {class_node.name}")
                
            except Exception as e:
                error_count += 1
                logger.error(f"  [ERROR] Error processing {java_file}: {e}")
                continue
        
        db.close()
        click.echo(f"Update complete. Processed: {processed_count}, Errors: {error_count}")
        logger.info("====== analyze 작업 완료 ======")
        
        return None, None
        
    except Exception as e:
        click.echo(f"Error updating classes: {e}")
        click.echo("Use --dry-run flag to parse without database connection.")
        exit(1)

@cli.command()
@click.option('--neo4j-uri', default=os.getenv("NEO4J_URI", "bolt://localhost:7687"), help='Neo4j URI')
@click.option('--neo4j-user', default=os.getenv("NEO4J_USER", "neo4j"), help='Neo4j username')
@click.option('--neo4j-password', default=os.getenv("NEO4J_PASSWORD"), help='Neo4j password')
@click.option('--neo4j-database', default=os.getenv("NEO4J_DATABASE", "neo4j"), help='Neo4j database name')
@click.option('--query', help='Custom Cypher query to execute')
@click.option('--basic', is_flag=True, help='Run basic class query')
@click.option('--detailed', is_flag=True, help='Run detailed class query with methods and properties')
@click.option('--inheritance', is_flag=True, help='Run inheritance relationship query')
@click.option('--package', is_flag=True, help='Run package-based class query')
def query(neo4j_uri, neo4j_user, neo4j_password, neo4j_database, query, basic, detailed, inheritance, package):
    """Execute queries against the Neo4j database."""
    
    # 미리 정의된 쿼리들
    queries = {
        'basic': """
        MATCH (c:Class)
        RETURN 
            c.name AS name,
            c.logical_name AS logical_name,
            c.file_path AS file_path,
            c.type AS type,
            c.source AS source
        ORDER BY c.name
        """,
        'detailed': """
        MATCH (c:Class)
        OPTIONAL MATCH (c)-[:HAS_METHOD]->(m:Method)
        OPTIONAL MATCH (c)-[:HAS_FIELD]->(p:Field)
        OPTIONAL MATCH (pkg:Package)-[:CONTAINS]->(c)
        RETURN 
            c.name AS class_name,
            c.logical_name AS class_logical_name,
            c.file_path AS file_path,
            c.type AS class_type,
            pkg.name AS package_name,
            collect(DISTINCT m.name) AS methods,
            collect(DISTINCT p.name) AS properties
        ORDER BY c.name
        """,
        'inheritance': """
        MATCH (c:Class)
        OPTIONAL MATCH (c)-[:EXTENDS]->(super:Class)
        OPTIONAL MATCH (c)-[:IMPLEMENTS]->(impl:Class)
        RETURN 
            c.name AS class_name,
            c.logical_name AS class_logical_name,
            c.type AS class_type,
            collect(DISTINCT super.name) AS extends,
            collect(DISTINCT impl.name) AS implements
        ORDER BY c.name
        """,
        'package': """
        MATCH (pkg:Package)-[:CONTAINS]->(c:Class)
        OPTIONAL MATCH (c)-[:HAS_METHOD]->(m:Method)
        OPTIONAL MATCH (c)-[:HAS_FIELD]->(p:Field)
        RETURN 
            pkg.name AS package_name,
            pkg.logical_name AS package_logical_name,
            collect(DISTINCT c.name) AS classes,
            count(DISTINCT m) AS total_methods,
            count(DISTINCT p) AS total_properties
        ORDER BY pkg.name
        """
    }
    
    # 실행할 쿼리 결정
    if query:
        cypher_query = query
        description = "Custom Query"
    elif basic:
        cypher_query = queries['basic']
        description = "Basic Class Query"
    elif detailed:
        cypher_query = queries['detailed']
        description = "Detailed Class Query"
    elif inheritance:
        cypher_query = queries['inheritance']
        description = "Inheritance Query"
    elif package:
        cypher_query = queries['package']
        description = "Package Query"
    else:
        click.echo("Error: Please specify a query type or provide a custom query.")
        click.echo("Available options: --basic, --detailed, --inheritance, --package, or --query")
        return
    
    try:
        # Connection Pool 초기화
        pool = get_connection_pool()
        if not pool.is_initialized():
            pool_size = int(os.getenv('NEO4J_POOL_SIZE', '10'))
            pool.initialize(neo4j_uri, neo4j_user, neo4j_password, neo4j_database, pool_size)
        
        conn = pool.acquire()
        try:
            with conn.session() as session:
                click.echo(f"Executing: {description}")
                click.echo("=" * 50)
                
                result = session.run(cypher_query)
                records = list(result)
                
                if not records:
                    click.echo("No results found.")
                    return
                
                # 첫 번째 레코드의 키들을 헤더로 사용
                headers = list(records[0].keys())
                
                # 헤더 출력
                click.echo(" | ".join(f"{header:20}" for header in headers))
                click.echo("-" * (len(headers) * 23))
                
                # 데이터 출력
                for record in records:
                    row = []
                    for header in headers:
                        value = record[header]
                        if value is None:
                            row.append("None")
                        elif isinstance(value, (list, dict)):
                            row.append(str(value)[:50] + "..." if len(str(value)) > 50 else str(value))
                        else:
                            row.append(str(value)[:20])
                    click.echo(" | ".join(f"{cell:20}" for cell in row))
                
                click.echo(f"\nTotal: {len(records)} results found.")
        finally:
            pool.release(conn)
            
    except Exception as e:
        click.echo(f"Error executing query: {e}")

@cli.command()
@click.option('--neo4j-uri', default=os.getenv("NEO4J_URI", "bolt://localhost:7687"), help='Neo4j URI')
@click.option('--neo4j-user', default=os.getenv("NEO4J_USER", "neo4j"), help='Neo4j username')
@click.option('--neo4j-database', default=os.getenv("NEO4J_DATABASE", "neo4j"), help='Neo4j database name')
@click.option('--class-name', required=True, help='Name of the class to analyze')
@click.option('--method-name', help='Specific method to analyze (optional)')
@click.option('--max-depth', default=10, help='Maximum depth of call chain to follow (default: 10)')
@click.option('--include-external', is_flag=True, help='Include calls to external libraries')
@click.option('--project-name', help='Project name for database analysis (optional, will auto-detect if not provided)')
@click.option('--image-format', default='none', type=click.Choice(['none', 'png', 'svg', 'pdf']), help='Image format (default: none - no image generation)')
@click.option('--image-width', default=1200, help='Image width in pixels (default: 1200)')
@click.option('--image-height', default=800, help='Image height in pixels (default: 800)')
@click.option('--format', default='plantuml', type=click.Choice(['mermaid', 'plantuml']), help='Diagram format (default: plantuml)')
@click.option('--output-dir', default=os.getenv("SEQUENCE_DIAGRAM_OUTPUT_DIR", "output/sequence-diagram"), help='Output directory for sequence diagrams (default: output/sequence-diagram)')
def sequence(neo4j_uri, neo4j_user, neo4j_database, class_name, method_name, max_depth, include_external, project_name, image_format, image_width, image_height, format, output_dir):
    """Generate sequence diagram for a specific class and optionally a method."""
    
    # 1. start() 호출
    context = start('sequence')
    
    result = {
        'success': False,
        'message': '',
        'stats': {},
        'error': None,
        'files': []
    }
    
    try:
        neo4j_password = os.getenv("NEO4J_PASSWORD")
        if not neo4j_password:
            result['error'] = "NEO4J_PASSWORD environment variable is not set"
            click.echo("Error: NEO4J_PASSWORD environment variable is not set.")
            click.echo("Please set NEO4J_PASSWORD in your .env file or environment variables.")
            return result
        
        click.echo(f"Connecting to Neo4j at {neo4j_uri} (database: {neo4j_database})...")
        
        # Connection Pool 초기화
        pool = get_connection_pool()
        if not pool.is_initialized():
            pool_size = int(os.getenv('NEO4J_POOL_SIZE', '10'))
            pool.initialize(neo4j_uri, neo4j_user, neo4j_password, neo4j_database, pool_size)
        
        conn = pool.acquire()
        generated_files = []  # 생성된 파일 목록 초기화
        try:
            # Create generator with specified format
            generator = SequenceDiagramGenerator(conn.driver, format=format)
            
            # Generate the sequence diagram
            click.echo(f"Generating {format} sequence diagram for class: {class_name}")
            if method_name:
                click.echo(f"Focusing on method: {method_name}")
            if project_name:
                click.echo(f"Using project: {project_name}")
            else:
                click.echo("Auto-detecting project name...")
            
            diagram = generator.generate_sequence_diagram(
                class_name=class_name,
                method_name=method_name,
                max_depth=max_depth,
                include_external_calls=include_external,
                project_name=project_name,
                output_dir=output_dir,
                image_format=image_format,
                image_width=image_width,
                image_height=image_height
            )
            
            # Check if diagram contains error message
            if isinstance(diagram, str) and diagram.startswith("Error:"):
                click.echo(f"Error: {diagram}")
                return
            
            # Handle different return types
            if isinstance(diagram, dict):
                if diagram.get("type") == "class":
                    click.echo(f"Generated {len(diagram['files'])} sequence diagram files for class '{class_name}':")
                    click.echo("")
                    
                    for file_info in diagram['files']:
                        click.echo(f"- Diagram: {os.path.basename(file_info['diagram_path'])}")
                        if file_info['image_path']:
                            click.echo(f"  Image: {os.path.basename(file_info['image_path'])}")
                    
                    click.echo(f"\nFiles saved in: {diagram['output_dir']}/ directory")
                    
                    # 파일 목록을 generated_files에 추가
                    for file_info in diagram['files']:
                        generated_files.append(file_info['diagram_path'])
                        if file_info['image_path']:
                            generated_files.append(file_info['image_path'])
                
                elif diagram.get("type") == "method":
                    click.echo(f"Generated sequence diagram for method '{method_name}':")
                    click.echo(f"- Diagram: {os.path.basename(diagram['diagram_path'])}")
                    if diagram['image_path']:
                        click.echo(f"- Image: {os.path.basename(diagram['image_path'])}")
                    
                    # 파일 목록을 generated_files에 추가
                    generated_files.append(diagram['diagram_path'])
                    if diagram['image_path']:
                        generated_files.append(diagram['image_path'])
            else:
                # Fallback for string return (should not happen with new implementation)
                click.echo(f"Diagram generated (length: {len(diagram)})")
                click.echo(diagram)
            
            # 결과 설정
            result['success'] = True
            result['message'] = 'Sequence diagram generated successfully'
            result['files'] = generated_files
            result['stats'] = {
                'class_name': class_name,
                'method_name': method_name,
                'max_depth': max_depth,
                'include_external': include_external,
                'format': format,
                'files_generated': len(generated_files)
            }
            
        except Exception as e:
            result['error'] = str(e)
            click.echo(f"Error generating sequence diagram: {e}")
            import traceback
            click.echo(f"Traceback: {traceback.format_exc()}")
    
    finally:
        # 4. end() 호출
        end(context, result)
    
    return result

@cli.command()
@click.option('--neo4j-uri', default=os.getenv("NEO4J_URI", "bolt://localhost:7687"), help='Neo4j URI')
@click.option('--neo4j-user', default=os.getenv("NEO4J_USER", "neo4j"), help='Neo4j username')
@click.option('--neo4j-database', default=os.getenv("NEO4J_DATABASE", "neo4j"), help='Neo4j database name')
def list_classes(neo4j_uri, neo4j_user, neo4j_database):
    """List all available classes in the database."""
    
    try:
        neo4j_password = os.getenv("NEO4J_PASSWORD")
        if not neo4j_password:
            click.echo("Error: NEO4J_PASSWORD environment variable is not set.")
            click.echo("Please set NEO4J_PASSWORD in your .env file or environment variables.")
            return
        
        driver = GraphDatabase.driver(neo4j_uri, auth=(neo4j_user, neo4j_password))
        generator = SequenceDiagramGenerator(driver)
        
        classes = generator.get_available_classes()
        
        if not classes:
            click.echo("No classes found in the database.")
            return
        
        click.echo("Available classes:")
        click.echo("=" * 80)
        click.echo(f"{'Class Name':<30} {'Package':<30} {'Type':<10}")
        click.echo("-" * 80)
        
        for cls in classes:
            package_name = cls.get('package_name') or 'N/A'
            class_type = cls.get('type') or 'N/A'
            click.echo(f"{cls['name']:<30} {package_name:<30} {class_type:<10}")
        
        click.echo(f"\nTotal: {len(classes)} classes found.")
        
    except Exception as e:
        click.echo(f"Error listing classes: {e}")
        return
    finally:
        driver.close()

@cli.command()
@click.option('--neo4j-uri', default=os.getenv("NEO4J_URI", "bolt://localhost:7687"), help='Neo4j URI')
@click.option('--neo4j-user', default=os.getenv("NEO4J_USER", "neo4j"), help='Neo4j username')
@click.option('--class-name', required=True, help='Name of the class to list methods for')
def list_methods(neo4j_uri, neo4j_user, class_name):
    """List all methods for a specific class."""
    
    try:
        neo4j_password = os.getenv("NEO4J_PASSWORD")
        if not neo4j_password:
            click.echo("Error: NEO4J_PASSWORD environment variable is not set.")
            click.echo("Please set NEO4J_PASSWORD in your .env file or environment variables.")
            return
        
        driver = GraphDatabase.driver(neo4j_uri, auth=(neo4j_user, neo4j_password))
        generator = SequenceDiagramGenerator(driver)
        
        methods = generator.get_class_methods(class_name)
        
        if not methods:
            click.echo(f"No methods found for class '{class_name}'.")
            return
        
        click.echo(f"Methods for class '{class_name}':")
        click.echo("=" * 80)
        click.echo(f"{'Method Name':<30} {'Return Type':<20} {'Logical Name':<30}")
        click.echo("-" * 80)
        
        for method in methods:
            click.echo(f"{method['name']:<30} {method['return_type']:<20} {method['logical_name']:<30}")
        
        click.echo(f"\nTotal: {len(methods)} methods found.")
        
    except Exception as e:
        click.echo(f"Error listing methods: {e}")
        return
    finally:
        driver.close()

@cli.command()
@click.option('--neo4j-uri', default=os.getenv("NEO4J_URI", "bolt://localhost:7687"), help='Neo4j URI')
@click.option('--neo4j-user', default=os.getenv("NEO4J_USER", "neo4j"), help='Neo4j username')
@click.option('--project-name', help='Project name to filter by (optional)')
@click.option('--output-format', default='excel', type=click.Choice(['excel', 'svg', 'png'], case_sensitive=False), 
              help='Output format: excel (*.xlsx), svg (*.svg), or png (*.png) (default: excel)')
@click.option('--auto-create-relationships', is_flag=True, default=True, help='Automatically create Method-SqlStatement relationships if needed (default: True)')
def crud_matrix(neo4j_uri, neo4j_user, project_name, output_format, auto_create_relationships):
    """Show CRUD matrix for classes and tables."""
    
    # 1. start() 호출
    context = start('crud-matrix')
    
    result = {
        'success': False,
        'message': '',
        'stats': {},
        'error': None,
        'files': []
    }
    
    try:
        neo4j_password = os.getenv("NEO4J_PASSWORD")
        if not neo4j_password:
            result['error'] = "NEO4J_PASSWORD environment variable is not set"
            click.echo("Error: NEO4J_PASSWORD environment variable is not set.")
            click.echo("Please set NEO4J_PASSWORD in your .env file or environment variables.")
            return result
        
        neo4j_uri = os.getenv("NEO4J_URI", "bolt://localhost:7687")
        neo4j_user = os.getenv("NEO4J_USER", "neo4j")
        neo4j_database = os.getenv("NEO4J_DATABASE", "neo4j")
        db = GraphDB(neo4j_uri, neo4j_user, neo4j_password, neo4j_database)
        
        click.echo("CRUD Matrix - Class to Table Operations")
        click.echo("=" * 80)
        
        # 먼저 CRUD 매트릭스 확인
        matrix = db.get_crud_matrix(project_name)
        
        # CRUD 매트릭스가 없고 자동 생성 옵션이 활성화된 경우 관계 생성
        if not matrix and auto_create_relationships:
            click.echo("No CRUD operations found. Creating Method-SqlStatement relationships...")
            relationships_created = db.create_method_sql_relationships(project_name)
            if relationships_created:
                click.echo(f"Created {relationships_created} Method-SqlStatement relationships.")
                # 관계 생성 후 다시 CRUD 매트릭스 확인
                matrix = db.get_crud_matrix(project_name)
            else:
                click.echo("No relationships could be created.")
        
        if not matrix:
            result['error'] = "No CRUD operations found"
            click.echo("No CRUD operations found.")
            if not auto_create_relationships:
                click.echo("Tip: Use --auto-create-relationships flag to automatically create Method-SqlStatement relationships.")
            return result
        
        click.echo(f"{'Package':<35} {'Class Name':<30} {'Method':<25} {'Schema':<10} {'Table':<20} {'Operations':<15}")
        click.echo("-" * 135)
        
        for row in matrix:
            package_name = row['package_name'] or 'N/A'
            class_name = row['class_name']
            method_name = row['method_name']
            schema = row['schema'] or 'unknown'
            table_name = row['table_name']
            operations = ', '.join(row['operations']) if row['operations'] else 'None'
            
            click.echo(f"{package_name:<35} {class_name:<30} {method_name:<25} {schema:<10} {table_name:<20} {operations:<15}")
        
        click.echo(f"\nTotal: {len(matrix)} class-table relationships.")
        
        # 출력 디렉토리 및 타임스탬프 준비
        output_dir = os.getenv("CRUD_MATRIX_OUTPUT_DIR", "./output/crud-matrix")
        os.makedirs(output_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        
        # 마크다운 파일 생성 (항상 실행)
        md_filename = f"CRUD_{project_name}_{timestamp}.md"
        md_filepath = os.path.join(output_dir, md_filename)
        
        # 마크다운 내용 생성
        markdown_content = f"# CRUD Matrix [Project : {project_name}]\n\n"
        markdown_content += f"Generated at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n"
        markdown_content += "| Package | Class Name | Method | Schema | Table | Operations |\n"
        markdown_content += "|---------|------------|--------|--------|-------|------------|\n"
        
        for row in matrix:
            package_name = row['package_name'] or 'N/A'
            class_name = row['class_name']
            method_name = row['method_name']
            schema = row['schema'] or 'unknown'
            table_name = row['table_name']
            operations = ', '.join(row['operations']) if row['operations'] else 'None'
            
            markdown_content += f"| {package_name} | {class_name} | {method_name} | {schema} | {table_name} | {operations} |\n"
        
        markdown_content += f"\n**Total:** {len(matrix)} class-table relationships.\n"
        
        # 마크다운 파일 저장
        with open(md_filepath, 'w', encoding='utf-8') as f:
            f.write(markdown_content)
        
        click.echo(f"\nCRUD matrix (Markdown) saved to: {md_filepath}")
        result['files'].append(md_filepath)
        
        # 추가 형식 파일 생성 (output_format이 지정된 경우)
        if output_format:
            if output_format.lower() == 'excel':
                excel_filename = f"CRUD_{project_name}_{timestamp}.xlsx"
                excel_filepath = os.path.join(output_dir, excel_filename)
                success = _save_crud_matrix_as_excel(matrix, project_name, excel_filepath)
                if success:
                    click.echo(f"CRUD matrix (Excel) saved to: {excel_filepath}")
                    result['files'].append(excel_filepath)
                else:
                    click.echo("Failed to generate Excel file. Check logs for details.")
            
            elif output_format.lower() in ['svg', 'png']:
                image_filename = f"CRUD_{project_name}_{timestamp}.{output_format.lower()}"
                image_filepath = os.path.join(output_dir, image_filename)
                success = _save_crud_matrix_as_image(matrix, project_name, image_filepath, output_format.lower())
                if success:
                    click.echo(f"CRUD matrix ({output_format.upper()}) saved to: {image_filepath}")
                    result['files'].append(image_filepath)
                else:
                    click.echo(f"Failed to generate {output_format.upper()} file. Check logs for details.")
        
        # 결과 설정
        result['success'] = True
        result['message'] = 'CRUD matrix generated successfully'
        
        # 테이블별로 그룹화하여 통계 계산
        table_groups = {}
        for row in matrix:
            table_name = row['table_name']
            if table_name not in table_groups:
                table_groups[table_name] = []
            table_groups[table_name].append(row)
        
        result['stats'] = {
            'total_relationships': len(matrix),
            'total_tables': len(table_groups),
            'output_format': output_format,
            'files_generated': len(result['files'])
        }
        
    except Exception as e:
        result['error'] = str(e)
        click.echo(f"Error getting CRUD matrix: {e}")
    
    finally:
        # 4. end() 호출
        end(context, result)
    
    return result

@cli.command()
@click.option('--neo4j-uri', default=os.getenv("NEO4J_URI", "bolt://localhost:7687"), help='Neo4j URI')
@click.option('--neo4j-user', default=os.getenv("NEO4J_USER", "neo4j"), help='Neo4j username')
@click.option('--project-name', help='Project name to filter by (optional)')
@click.option('--auto-create-relationships', is_flag=True, default=True, help='Automatically create Method-SqlStatement relationships if needed (default: True)')
def db_analysis(neo4j_uri, neo4j_user, project_name, auto_create_relationships):
    """Show database call relationship analysis."""
    
    try:
        neo4j_password = os.getenv("NEO4J_PASSWORD")
        if not neo4j_password:
            click.echo("Error: NEO4J_PASSWORD environment variable is not set.")
            click.echo("Please set NEO4J_PASSWORD in your .env file or environment variables.")
            return
        
        driver = GraphDatabase.driver(neo4j_uri, auth=(neo4j_user, neo4j_password))
        db = GraphDB(neo4j_uri, neo4j_user, neo4j_password, neo4j_database)
        
        click.echo("Database Call Relationship Analysis")
        click.echo("=" * 80)
        
        # 먼저 SQL 문 통계 확인
        sql_stats = db.get_sql_statistics(project_name)
        
        # 통계가 없고 자동 생성 옵션이 활성화된 경우 관계 생성
        if not sql_stats and auto_create_relationships:
            click.echo("No SQL statistics found. Creating Method-SqlStatement relationships...")
            relationships_created = db.create_method_sql_relationships(project_name)
            if relationships_created:
                click.echo(f"Created {relationships_created} Method-SqlStatement relationships.")
                # 관계 생성 후 다시 SQL 통계 확인
                sql_stats = db.get_sql_statistics(project_name)
            else:
                click.echo("No relationships could be created.")
        if sql_stats:
            click.echo(f"\nSQL Statistics:")
            click.echo(f"  Total SQL statements: {sql_stats['total_sql']}")
            click.echo(f"  SELECT statements: {sql_stats.get('SELECT', 0)}")
            click.echo(f"  INSERT statements: {sql_stats.get('INSERT', 0)}")
            click.echo(f"  UPDATE statements: {sql_stats.get('UPDATE', 0)}")
            click.echo(f"  DELETE statements: {sql_stats.get('DELETE', 0)}")
        
        # 테이블 사용 통계
        table_stats = db.get_table_usage_statistics(project_name)
        if table_stats:
            click.echo(f"\nTable Usage Statistics:")
            click.echo(f"{'Table Name':<30} {'Access Count':<15} {'Operations':<20}")
            click.echo("-" * 65)
            for table in table_stats:
                table_name = table['table_name']
                access_count = table['access_count']
                operations = ', '.join(table['operations'])
                click.echo(f"{table_name:<30} {access_count:<15} {operations:<20}")
        
        # 복잡도 분석
        complexity_stats = db.get_sql_complexity_statistics(project_name)
        if complexity_stats:
            click.echo(f"\nSQL Complexity Analysis:")
            click.echo(f"  Simple queries: {complexity_stats.get('simple', 0)}")
            click.echo(f"  Medium queries: {complexity_stats.get('medium', 0)}")
            click.echo(f"  Complex queries: {complexity_stats.get('complex', 0)}")
            click.echo(f"  Very complex queries: {complexity_stats.get('very_complex', 0)}")
        
        # 매퍼별 SQL 분포
        mapper_stats = db.get_mapper_sql_distribution(project_name)
        if mapper_stats:
            click.echo(f"\nMapper SQL Distribution:")
            click.echo(f"{'Mapper Name':<30} {'SQL Count':<15} {'SQL Types':<20}")
            click.echo("-" * 65)
            for mapper in mapper_stats:
                mapper_name = mapper['mapper_name']
                sql_count = mapper['sql_count']
                sql_types = ', '.join(mapper['sql_types'])
                click.echo(f"{mapper_name:<30} {sql_count:<15} {sql_types:<20}")
        
    except Exception as e:
        click.echo(f"Error getting database analysis: {e}")
        return
    finally:
        driver.close()

@cli.command()
@click.option('--neo4j-uri', default=os.getenv("NEO4J_URI", "bolt://localhost:7687"), help='Neo4j URI')
@click.option('--neo4j-user', default=os.getenv("NEO4J_USER", "neo4j"), help='Neo4j username')
@click.option('--project-name', help='Project name to filter by (optional)')
@click.option('--auto-create-relationships', is_flag=True, default=True, help='Automatically create Method-SqlStatement relationships if needed (default: True)')
def table_summary(neo4j_uri, neo4j_user, project_name, auto_create_relationships):
    """Show CRUD summary for each table."""
    
    try:
        neo4j_password = os.getenv("NEO4J_PASSWORD")
        if not neo4j_password:
            click.echo("Error: NEO4J_PASSWORD environment variable is not set.")
            click.echo("Please set NEO4J_PASSWORD in your .env file or environment variables.")
            return
        
        driver = GraphDatabase.driver(neo4j_uri, auth=(neo4j_user, neo4j_password))
        db = GraphDB(neo4j_uri, neo4j_user, neo4j_password, neo4j_database)
        
        click.echo("Table CRUD Summary")
        click.echo("=" * 60)
        
        # 먼저 테이블 요약 확인
        summary = db.get_table_crud_summary(project_name)
        
        # 요약이 없고 자동 생성 옵션이 활성화된 경우 관계 생성
        if not summary and auto_create_relationships:
            click.echo("No table CRUD operations found. Creating Method-SqlStatement relationships...")
            relationships_created = db.create_method_sql_relationships(project_name)
            if relationships_created:
                click.echo(f"Created {relationships_created} Method-SqlStatement relationships.")
                # 관계 생성 후 다시 테이블 요약 확인
                summary = db.get_table_crud_summary(project_name)
            else:
                click.echo("No relationships could be created.")
        
        if not summary:
            click.echo("No tables found.")
            return
        
        for row in summary:
            table_name = row['table_name']
            operations = row['operations']
            
            click.echo(f"\nTable: {table_name}")
            click.echo("-" * 40)
            
            for op in operations:
                operation = op['operation']
                count = op['count']
                click.echo(f"  {operation}: {count} statements")
        
        click.echo(f"\nTotal: {len(summary)} tables.")
        
    except Exception as e:
        click.echo(f"Error getting table summary: {e}")
        return
    finally:
        driver.close()

@cli.command()
@click.option('--neo4j-uri', default=os.getenv("NEO4J_URI", "bolt://localhost:7687"), help='Neo4j URI')
@click.option('--neo4j-user', default=os.getenv("NEO4J_USER", "neo4j"), help='Neo4j username')
@click.option('--project-name', required=True, help='Project name to analyze')
@click.option('--start-class', help='Starting class for call chain analysis (optional)')
@click.option('--start-method', help='Starting method for call chain analysis (optional)')
@click.option('--output-file', help='Output file to save the analysis results (optional)')
@click.option('--auto-create-relationships', is_flag=True, default=True, help='Automatically create Method-SqlStatement relationships if needed (default: True)')
def db_call_chain(neo4j_uri, neo4j_user, project_name, start_class, start_method, output_file, auto_create_relationships):
    """Analyze database call chain relationships."""
    
    try:
        neo4j_password = os.getenv("NEO4J_PASSWORD")
        if not neo4j_password:
            click.echo("Error: NEO4J_PASSWORD environment variable is not set.")
            click.echo("Please set NEO4J_PASSWORD in your .env file or environment variables.")
            return
        
        driver = GraphDatabase.driver(neo4j_uri, auth=(neo4j_user, neo4j_password))
        analysis_service = DBCallAnalysisService(driver)
        
        click.echo("Database Call Chain Analysis")
        click.echo("=" * 50)
        
        if start_class and start_method:
            click.echo(f"Analyzing call chain from {start_class}.{start_method}")
        elif start_class:
            click.echo(f"Analyzing call chain from class {start_class}")
        else:
            click.echo(f"Analyzing call chain for project {project_name}")
        
        # 먼저 분석 실행
        result = analysis_service.analyze_call_chain(project_name, start_class, start_method)
        
        # 분석 결과가 없고 자동 생성 옵션이 활성화된 경우 관계 생성
        if auto_create_relationships and ('error' in result or not result.get('call_chain')):
            click.echo("No call chain analysis found. Creating Method-SqlStatement relationships...")
            graph_db = GraphDB(neo4j_uri, neo4j_user, neo4j_password, neo4j_database)
            relationships_created = graph_db.create_method_sql_relationships(project_name)
            if relationships_created:
                click.echo(f"Created {relationships_created} Method-SqlStatement relationships.")
                # 관계 생성 후 다시 분석 실행
                result = analysis_service.analyze_call_chain(project_name, start_class, start_method)
            else:
                click.echo("No relationships could be created.")
        
        if 'error' in result:
            click.echo(f"Error: {result['error']}")
            return
        
        # 결과 출력
        call_chain = result['call_chain']
        missing_nodes = result['missing_nodes']
        summary = result['analysis_summary']
        
        click.echo(f"\nAnalysis Summary:")
        click.echo(f"  Total calls: {summary['total_calls']}")
        click.echo(f"  Unique classes: {summary['unique_classes']}")
        click.echo(f"  Unique methods: {summary['unique_methods']}")
        click.echo(f"  Unique SQL statements: {summary['unique_sql_statements']}")
        click.echo(f"  Unique tables: {summary['unique_tables']}")
        click.echo(f"  Unique columns: {summary['unique_columns']}")
        click.echo(f"  Missing tables: {summary['missing_tables_count']}")
        click.echo(f"  Missing columns: {summary['missing_columns_count']}")
        
        if missing_nodes['missing_tables']:
            click.echo(f"\nMissing Tables (❌):")
            for table in missing_nodes['missing_tables']:
                click.echo(f"  - {table}")
        
        if missing_nodes['missing_columns']:
            click.echo(f"\nMissing Columns (❌):")
            for column in missing_nodes['missing_columns']:
                click.echo(f"  - {column}")
        
        # 호출 체인 상세 정보
        if call_chain:
            click.echo(f"\nCall Chain Details:")
            click.echo("-" * 80)
            click.echo(f"{'Source':<25} {'Target':<25} {'SQL Type':<10} {'Table':<20}")
            click.echo("-" * 80)
            
            for call in call_chain[:20]:  # 처음 20개만 표시
                source = f"{call['source_class']}.{call['source_method']}" if call['source_method'] else call['source_class']
                target = f"{call['target_class']}.{call['target_method']}" if call['target_method'] else call['target_class']
                sql_type = call['sql_type'] or 'N/A'
                table = call['table_name'] or 'N/A'
                
                click.echo(f"{source:<25} {target:<25} {sql_type:<10} {table:<20}")
            
            if len(call_chain) > 20:
                click.echo(f"... and {len(call_chain) - 20} more calls")
        
        # 파일로 저장
        if output_file:
            import json
            with open(output_file, 'w', encoding='utf-8') as f:
                json.dump(result, f, indent=2, ensure_ascii=False)
            click.echo(f"\nAnalysis results saved to: {output_file}")
        
    except Exception as e:
        click.echo(f"Error analyzing call chain: {e}")
        return
    finally:
        driver.close()

@cli.command()
@click.option('--neo4j-uri', default=os.getenv("NEO4J_URI", "bolt://localhost:7687"), help='Neo4j URI')
@click.option('--neo4j-user', default=os.getenv("NEO4J_USER", "neo4j"), help='Neo4j username')
@click.option('--project-name', required=True, help='Project name to analyze')
@click.option('--output-file', help='Output file to save the CRUD matrix (optional)')
@click.option('--output-excel', help='Output Excel file to save the CRUD matrix (optional)')
@click.option('--auto-create-relationships', is_flag=True, default=True, help='Automatically create Method-SqlStatement relationships if needed (default: True)')
def crud_analysis(neo4j_uri, neo4j_user, project_name, output_file, output_excel, auto_create_relationships):
    """Generate CRUD matrix analysis."""
    
    try:
        neo4j_password = os.getenv("NEO4J_PASSWORD")
        if not neo4j_password:
            click.echo("Error: NEO4J_PASSWORD environment variable is not set.")
            click.echo("Please set NEO4J_PASSWORD in your .env file or environment variables.")
            return
        
        driver = GraphDatabase.driver(neo4j_uri, auth=(neo4j_user, neo4j_password))
        analysis_service = DBCallAnalysisService(driver)
        
        click.echo("CRUD Matrix Analysis (SQL 호출 클래스만)")
        click.echo("=" * 50)
        
        # 먼저 CRUD 매트릭스 생성 시도
        result = analysis_service.generate_crud_table_matrix(project_name)
        
        # CRUD 매트릭스가 없고 자동 생성 옵션이 활성화된 경우 관계 생성
        if auto_create_relationships and ('error' in result or not result.get('table_matrix')):
            click.echo("No CRUD operations found. Creating Method-SqlStatement relationships...")
            graph_db = GraphDB(neo4j_uri, neo4j_user, neo4j_password, neo4j_database)
            relationships_created = graph_db.create_method_sql_relationships(project_name)
            if relationships_created:
                click.echo(f"Created {relationships_created} Method-SqlStatement relationships.")
                # 관계 생성 후 다시 CRUD 매트릭스 생성
                result = analysis_service.generate_crud_table_matrix(project_name)
            else:
                click.echo("No relationships could be created.")
        
        if 'error' in result:
            click.echo(f"Error: {result['error']}")
            return
        
        table_matrix = result['table_matrix']
        class_names = result['class_names']
        table_names = result['table_names']
        summary = result['summary']
        
        click.echo(f"\nCRUD Summary:")
        click.echo(f"  Total classes: {summary['total_classes']}")
        click.echo(f"  Total tables: {summary['total_tables']}")
        click.echo(f"  Create operations: {summary['crud_stats']['C']}")
        click.echo(f"  Read operations: {summary['crud_stats']['R']}")
        click.echo(f"  Update operations: {summary['crud_stats']['U']}")
        click.echo(f"  Delete operations: {summary['crud_stats']['D']}")
        click.echo(f"  Other operations: {summary['crud_stats']['O']}")
        
        if summary['most_active_class']:
            click.echo(f"  Most active class: {summary['most_active_class']}")
        if summary['most_used_table']:
            click.echo(f"  Most used table: {summary['most_used_table']}")
        
        # CRUD 매트릭스 표 출력
        if table_matrix and table_names:
            click.echo(f"\nCRUD Matrix (Class vs Table):")
            
            # 테이블 헤더 계산
            class_name_width = max(len("Class (Package)"), max(len(f"{row['class_name']} ({row['package_name']})") for row in table_matrix)) if table_matrix else 20
            table_width = 18  # 각 테이블 컬럼 너비 (스키마 정보 포함으로 더 넓게)
            
            # 헤더 출력
            header = f"{'Class (Package)':<{class_name_width}}"
            for table_name in table_names:
                # 테이블 이름을 12자로 제한
                short_name = table_name[:12] if len(table_name) > 12 else table_name
                header += f" {short_name:<{table_width}}"
            click.echo(header)
            
            # 구분선 출력
            separator = "-" * class_name_width
            for _ in table_names:
                separator += " " + "-" * table_width
            click.echo(separator)
            
            # 데이터 행 출력
            for row in table_matrix:
                class_name = row['class_name']
                package_name = row.get('package_name', 'N/A')
                class_display = f"{class_name} ({package_name})"
                line = f"{class_display:<{class_name_width}}"
                for table_name in table_names:
                    operations = row.get(table_name, '-')
                    line += f" {operations:<{table_width}}"
                click.echo(line)
        else:
            click.echo(f"\nSQL을 직접 호출하는 클래스가 없습니다.")
            click.echo(f"다음을 확인해주세요:")
            click.echo(f"  1. Java 객체 분석이 완료되었는지 확인")
            click.echo(f"  2. MyBatis Mapper와 SQL 문이 분석되었는지 확인")
            click.echo(f"  3. 프로젝트 이름이 올바른지 확인")
        
        # 파일로 저장
        if output_file:
            import json
            with open(output_file, 'w', encoding='utf-8') as f:
                json.dump(result, f, indent=2, ensure_ascii=False)
            click.echo(f"\nCRUD matrix saved to: {output_file}")
        
        # Excel 파일로 저장
        if output_excel:
            success = analysis_service.generate_crud_excel(project_name, output_excel)
            if success:
                click.echo(f"CRUD matrix Excel file saved to: {output_excel}")
            else:
                click.echo("Failed to generate Excel file. Check logs for details.")
        
    except Exception as e:
        click.echo(f"Error generating CRUD matrix: {e}")
        return
    finally:
        driver.close()

@cli.command()
@click.option('--neo4j-uri', default=os.getenv("NEO4J_URI", "bolt://localhost:7687"), help='Neo4j URI')
@click.option('--neo4j-user', default=os.getenv("NEO4J_USER", "neo4j"), help='Neo4j username')
@click.option('--project-name', required=True, help='Project name to analyze')
@click.option('--start-class', help='Starting class for diagram (optional)')
@click.option('--start-method', help='Starting method for diagram (optional)')
@click.option('--output-file', help='Output file to save the diagram (optional)')
@click.option('--output-image', help='Output image file (PNG/SVG/PDF) - requires mermaid-cli')
@click.option('--image-format', default='png', type=click.Choice(['png', 'svg', 'pdf']), help='Image format (default: png)')
@click.option('--image-width', default=1200, help='Image width in pixels (default: 1200)')
@click.option('--image-height', default=800, help='Image height in pixels (default: 800)')
@click.option('--auto-create-relationships', is_flag=True, default=True, help='Automatically create Method-SqlStatement relationships if needed (default: True)')
def db_call_diagram(neo4j_uri, neo4j_user, project_name, start_class, start_method, output_file, output_image, image_format, image_width, image_height, auto_create_relationships):
    """Generate database call chain diagram."""
    
    try:
        neo4j_password = os.getenv("NEO4J_PASSWORD")
        if not neo4j_password:
            click.echo("Error: NEO4J_PASSWORD environment variable is not set.")
            click.echo("Please set NEO4J_PASSWORD in your .env file or environment variables.")
            return
        
        driver = GraphDatabase.driver(neo4j_uri, auth=(neo4j_user, neo4j_password))
        analysis_service = DBCallAnalysisService(driver)
        
        click.echo("Database Call Chain Diagram")
        click.echo("=" * 50)
        
        if start_class and start_method:
            click.echo(f"Generating diagram from {start_class}.{start_method}")
        elif start_class:
            click.echo(f"Generating diagram from class {start_class}")
        else:
            click.echo(f"Generating diagram for project {project_name}")
        
        # 먼저 다이어그램 생성 시도
        diagram = analysis_service.generate_call_chain_diagram(project_name, start_class, start_method)
        
        # 다이어그램이 오류이고 자동 생성 옵션이 활성화된 경우 관계 생성
        if auto_create_relationships and diagram.startswith("오류:"):
            click.echo("No call chain diagram found. Creating Method-SqlStatement relationships...")
            graph_db = GraphDB(neo4j_uri, neo4j_user, neo4j_password, neo4j_database)
            relationships_created = graph_db.create_method_sql_relationships(project_name)
            if relationships_created:
                click.echo(f"Created {relationships_created} Method-SqlStatement relationships.")
                # 관계 생성 후 다시 다이어그램 생성
                diagram = analysis_service.generate_call_chain_diagram(project_name, start_class, start_method)
            else:
                click.echo("No relationships could be created.")
        
        if diagram.startswith("오류:"):
            click.echo(f"Error: {diagram}")
            return
        
        # 파일로 저장
        if output_file:
            with open(output_file, 'w', encoding='utf-8') as f:
                f.write(diagram)
            click.echo(f"Diagram saved to: {output_file}")
        else:
            # 기본 파일명으로 저장
            default_filename = f"db_call_chain_{project_name}.md"
            with open(default_filename, 'w', encoding='utf-8') as f:
                f.write(diagram)
            click.echo(f"Diagram saved to: {default_filename}")
        
        # 이미지로 변환
        if output_image:
            convert_to_image(diagram, output_image, image_format, image_width, image_height)
        
        # 다이어그램 미리보기
        click.echo("\n" + "="*50)
        click.echo("DATABASE CALL CHAIN DIAGRAM")
        click.echo("="*50)
        click.echo(diagram)
        click.echo("="*50)
        
    except Exception as e:
        click.echo(f"Error generating diagram: {e}")
        return
    finally:
        driver.close()

@cli.command()
@click.option('--neo4j-uri', default=os.getenv("NEO4J_URI", "bolt://localhost:7687"), help='Neo4j URI')
@click.option('--neo4j-user', default=os.getenv("NEO4J_USER", "neo4j"), help='Neo4j username')
@click.option('--project-name', required=True, help='Project name to analyze')
@click.option('--output-format', default='excel', type=click.Choice(['excel', 'svg', 'png']), help='Output format (default: excel)')
@click.option('--image-width', default=1200, help='Image width in pixels (default: 1200)')
@click.option('--image-height', default=800, help='Image height in pixels (default: 800)')
@click.option('--auto-create-relationships', is_flag=True, default=True, help='Automatically create Method-SqlStatement relationships if needed (default: True)')
def crud_visualization(neo4j_uri, neo4j_user, project_name, output_format, image_width, image_height, auto_create_relationships):
    """Generate CRUD matrix visualization diagram showing class-table relationships."""
    
    try:
        neo4j_password = os.getenv("NEO4J_PASSWORD")
        if not neo4j_password:
            click.echo("Error: NEO4J_PASSWORD environment variable is not set.")
            click.echo("Please set NEO4J_PASSWORD in your .env file or environment variables.")
            return
        
        driver = GraphDatabase.driver(neo4j_uri, auth=(neo4j_user, neo4j_password))
        analysis_service = DBCallAnalysisService(driver)
        
        click.echo("CRUD Matrix Visualization")
        click.echo("=" * 50)
        click.echo(f"Generating {output_format.upper()} visualization for project: {project_name}")
        
        # CRUD 매트릭스 데이터 가져오기
        result = analysis_service.generate_crud_matrix(project_name)
        
        # CRUD 매트릭스가 없고 자동 생성 옵션이 활성화된 경우 관계 생성
        if auto_create_relationships and ('error' in result or not result.get('class_matrix')):
            click.echo("No CRUD data found. Creating Method-SqlStatement relationships...")
            graph_db = GraphDB(neo4j_uri, neo4j_user, neo4j_password, neo4j_database)
            relationships_created = graph_db.create_method_sql_relationships(project_name)
            if relationships_created:
                click.echo(f"Created {relationships_created} Method-SqlStatement relationships.")
                # 관계 생성 후 다시 CRUD 매트릭스 생성
                result = analysis_service.generate_crud_matrix(project_name)
            else:
                click.echo("No relationships could be created.")
        
        if 'error' in result or not result.get('class_matrix'):
            click.echo(f"Error: {result.get('error', 'No CRUD data found')}")
            return
        
        # 출력 디렉토리 및 타임스탬프 준비
        output_dir = os.getenv("CRUD_MATRIX_OUTPUT_DIR", "./output/crud-matrix")
        os.makedirs(output_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        
        # output_format에 따라 파일 저장
        if output_format == 'excel':
            # Excel 파일로 저장
            excel_filename = f"CRUD_visualization_{project_name}_{timestamp}.xlsx"
            excel_filepath = os.path.join(output_dir, excel_filename)
            success = _save_crud_matrix_as_excel(result, project_name, excel_filepath)
            if success:
                click.echo(f"CRUD visualization (Excel) saved to: {excel_filepath}")
            else:
                click.echo("Failed to save Excel file")
        else:
            # SVG/PNG 이미지로 저장
            image_filename = f"CRUD_visualization_{project_name}_{timestamp}.{output_format.lower()}"
            image_filepath = os.path.join(output_dir, image_filename)
            success = _save_crud_matrix_as_image(result, project_name, image_filepath, output_format.lower())
            if success:
                click.echo(f"CRUD visualization ({output_format.upper()}) saved to: {image_filepath}")
            else:
                click.echo(f"Failed to save {output_format.upper()} file")
        
        # 요약 정보 표시
        class_matrix = result.get('class_matrix', [])
        table_matrix = result.get('table_matrix', [])
        
        click.echo("\n" + "="*50)
        click.echo("CRUD MATRIX SUMMARY")
        click.echo("="*50)
        click.echo(f"Total classes: {len(class_matrix)}")
        click.echo(f"Total tables: {len(table_matrix)}")
        
        if class_matrix:
            click.echo("\nClasses with database operations:")
            for class_data in class_matrix[:10]:  # 최대 10개만 표시
                class_name = class_data.get('class_name', 'Unknown')
                tables = class_data.get('tables', [])
                table_count = len(tables) if isinstance(tables, list) else 0
                click.echo(f"  - {class_name}: {table_count} tables")
            
            if len(class_matrix) > 10:
                click.echo(f"  ... and {len(class_matrix) - 10} more classes")
        
        click.echo("="*50)
        
    except Exception as e:
        click.echo(f"Error generating diagram: {e}")
        return
    finally:
        driver.close()

@cli.command()
@click.option('--neo4j-uri', default=os.getenv("NEO4J_URI", "bolt://localhost:7687"), help='Neo4j URI')
@click.option('--neo4j-user', default=os.getenv("NEO4J_USER", "neo4j"), help='Neo4j username')
@click.option('--project-name', required=True, help='Project name to analyze')
@click.option('--table-name', required=True, help='Table name to analyze impact for')
@click.option('--output-file', help='Output file to save the impact analysis (optional)')
@click.option('--auto-create-relationships', is_flag=True, default=True, help='Automatically create Method-SqlStatement relationships if needed (default: True)')
def table_impact(neo4j_uri, neo4j_user, project_name, table_name, output_file, auto_create_relationships):
    """Analyze impact of table changes on application code."""
    
    try:
        neo4j_password = os.getenv("NEO4J_PASSWORD")
        if not neo4j_password:
            click.echo("Error: NEO4J_PASSWORD environment variable is not set.")
            click.echo("Please set NEO4J_PASSWORD in your .env file or environment variables.")
            return
        
        driver = GraphDatabase.driver(neo4j_uri, auth=(neo4j_user, neo4j_password))
        analysis_service = DBCallAnalysisService(driver)
        
        click.echo("Table Impact Analysis")
        click.echo("=" * 50)
        click.echo(f"Analyzing impact of changes to table: {table_name}")
        
        # 먼저 영향도 분석 시도
        result = analysis_service.analyze_table_impact(project_name, table_name)
        
        # 분석 결과가 없고 자동 생성 옵션이 활성화된 경우 관계 생성
        if auto_create_relationships and ('error' in result or not result.get('impacted_classes')):
            click.echo("No impact analysis found. Creating Method-SqlStatement relationships...")
            graph_db = GraphDB(neo4j_uri, neo4j_user, neo4j_password, neo4j_database)
            relationships_created = graph_db.create_method_sql_relationships(project_name)
            if relationships_created:
                click.echo(f"Created {relationships_created} Method-SqlStatement relationships.")
                # 관계 생성 후 다시 영향도 분석
                result = analysis_service.analyze_table_impact(project_name, table_name)
            else:
                click.echo("No relationships could be created.")
        
        if 'error' in result:
            click.echo(f"Error: {result['error']}")
            return
        
        impacted_classes = result['impacted_classes']
        summary = result['summary']
        
        click.echo(f"\nImpact Summary:")
        click.echo(f"  Table: {summary['table_name']}")
        click.echo(f"  Impacted classes: {summary['total_impacted_classes']}")
        click.echo(f"  Impacted methods: {summary['total_impacted_methods']}")
        click.echo(f"  SQL statements: {summary['total_sql_statements']}")
        click.echo(f"  CRUD operations: {', '.join(summary['crud_operations'])}")
        
        if summary['high_complexity_sql']:
            click.echo(f"  High complexity SQL: {len(summary['high_complexity_sql'])}")
        
        # 영향받는 클래스 상세 정보
        if impacted_classes:
            click.echo(f"\nImpacted Classes:")
            click.echo("-" * 80)
            click.echo(f"{'Class':<25} {'Method':<25} {'SQL Type':<10} {'Complexity':<12}")
            click.echo("-" * 80)
            
            for cls in impacted_classes:
                class_name = cls['class_name']
                method_name = cls['method_name'] or 'N/A'
                sql_type = cls['sql_type'] or 'N/A'
                complexity = str(cls['complexity_score']) if cls['complexity_score'] else 'N/A'
                
                click.echo(f"{class_name:<25} {method_name:<25} {sql_type:<10} {complexity:<12}")
        
        # 고복잡도 SQL 상세 정보
        if summary['high_complexity_sql']:
            click.echo(f"\nHigh Complexity SQL Statements:")
            click.echo("-" * 60)
            for sql in summary['high_complexity_sql']:
                click.echo(f"  {sql['class_name']}.{sql['method_name']} - {sql['sql_type']} (complexity: {sql['complexity_score']})")
        
        # 파일로 저장
        if output_file:
            import json
            with open(output_file, 'w', encoding='utf-8') as f:
                json.dump(result, f, indent=2, ensure_ascii=False)
            click.echo(f"\nImpact analysis saved to: {output_file}")
        
    except Exception as e:
        click.echo(f"Error analyzing table impact: {e}")
        return
    finally:
        driver.close()

@cli.command()
@click.option('--neo4j-uri', default=os.getenv("NEO4J_URI", "bolt://localhost:7687"), help='Neo4j URI')
@click.option('--neo4j-user', default=os.getenv("NEO4J_USER", "neo4j"), help='Neo4j username')
@click.option('--project-name', required=True, help='Project name to analyze')
@click.option('--output-file', help='Output file to save the statistics (optional)')
@click.option('--auto-create-relationships', is_flag=True, default=True, help='Automatically create Method-SqlStatement relationships if needed (default: True)')
def db_statistics(neo4j_uri, neo4j_user, project_name, output_file, auto_create_relationships):
    """Show database usage statistics."""
    
    try:
        neo4j_password = os.getenv("NEO4J_PASSWORD")
        if not neo4j_password:
            click.echo("Error: NEO4J_PASSWORD environment variable is not set.")
            click.echo("Please set NEO4J_PASSWORD in your .env file or environment variables.")
            return
        
        driver = GraphDatabase.driver(neo4j_uri, auth=(neo4j_user, neo4j_password))
        analysis_service = DBCallAnalysisService(driver)
        
        click.echo("Database Usage Statistics")
        click.echo("=" * 50)
        
        # 먼저 통계 조회 시도
        result = analysis_service.get_database_usage_statistics(project_name)
        
        # 통계가 없고 자동 생성 옵션이 활성화된 경우 관계 생성
        if auto_create_relationships and ('error' in result or not result.get('sql_statistics')):
            click.echo("No database statistics found. Creating Method-SqlStatement relationships...")
            graph_db = GraphDB(neo4j_uri, neo4j_user, neo4j_password, neo4j_database)
            relationships_created = graph_db.create_method_sql_relationships(project_name)
            if relationships_created:
                click.echo(f"Created {relationships_created} Method-SqlStatement relationships.")
                # 관계 생성 후 다시 통계 조회
                result = analysis_service.get_database_usage_statistics(project_name)
            else:
                click.echo("No relationships could be created.")
        
        if 'error' in result:
            click.echo(f"Error: {result['error']}")
            return
        
        sql_stats = result['sql_statistics']
        table_usage = result['table_usage']
        complexity_stats = result['complexity_statistics']
        
        # SQL 통계
        if sql_stats:
            click.echo(f"\nSQL Statistics:")
            click.echo(f"  Total SQL statements: {sql_stats['total_sql']}")
            click.echo(f"  SELECT statements: {sql_stats.get('SELECT', 0)}")
            click.echo(f"  INSERT statements: {sql_stats.get('INSERT', 0)}")
            click.echo(f"  UPDATE statements: {sql_stats.get('UPDATE', 0)}")
            click.echo(f"  DELETE statements: {sql_stats.get('DELETE', 0)}")
        
        # 테이블 사용 통계
        if table_usage:
            click.echo(f"\nTable Usage Statistics:")
            click.echo("-" * 60)
            click.echo(f"{'Table Name':<30} {'Access Count':<15} {'Operations':<20}")
            click.echo("-" * 60)
            
            for table in table_usage:
                table_name = table['table_name']
                access_count = table['access_count']
                operations = ', '.join(table['operations'])
                click.echo(f"{table_name:<30} {access_count:<15} {operations:<20}")
        
        # 복잡도 통계
        if complexity_stats:
            click.echo(f"\nSQL Complexity Statistics:")
            click.echo(f"  Simple queries: {complexity_stats.get('simple', 0)}")
            click.echo(f"  Medium queries: {complexity_stats.get('medium', 0)}")
            click.echo(f"  Complex queries: {complexity_stats.get('complex', 0)}")
            click.echo(f"  Very complex queries: {complexity_stats.get('very_complex', 0)}")
        
        # 파일로 저장
        if output_file:
            import json
            with open(output_file, 'w', encoding='utf-8') as f:
                json.dump(result, f, indent=2, ensure_ascii=False)
            click.echo(f"\nStatistics saved to: {output_file}")
        
    except Exception as e:
        click.echo(f"Error getting database statistics: {e}")
        return
    finally:
        driver.close()

@cli.command()
@click.option('--java-source-folder', help='Path to Java source folder (default: current directory)')
@click.option('--project-name', help='Project name (if not provided, will be extracted from folder name)')
@click.option('--db-script-folder', help='Path to database script folder')
@click.option('--neo4j-uri', default=os.getenv("NEO4J_URI", "bolt://localhost:7687"), help='Neo4j URI')
@click.option('--neo4j-user', default=os.getenv("NEO4J_USER", "neo4j"), help='Neo4j username')
@click.option('--neo4j-password', default=os.getenv("NEO4J_PASSWORD"), help='Neo4j password')
@click.option('--neo4j-database', default=os.getenv("NEO4J_DATABASE", "neo4j"), help='Neo4j database name')
@click.option('--clean', is_flag=True, help='Clean database before analysis')
@click.option('--dry-run', is_flag=True, help='Parse without database connection')
@click.option('--concurrent', is_flag=True, help='Use concurrent processing for Java analysis')
@click.option('--workers', type=int, help='Number of worker threads for concurrent processing')
@click.option('--java-object', is_flag=True, help='Analyze Java objects only')
@click.option('--db-object', is_flag=True, help='Analyze database objects only')
@click.option('--all-objects', is_flag=True, help='Analyze both Java and database objects')
@click.option('--class-name', help='Analyze specific class only')
@click.option('--update', is_flag=True, help='Update existing classes')
def analyze(java_source_folder, project_name, db_script_folder, neo4j_uri, neo4j_user, neo4j_password, neo4j_database, clean, dry_run, concurrent, workers, java_object, db_object, all_objects, class_name, update):
    """Analyze Java project and database objects."""
    # start() 함수로 공통 초기화
    context = start('analyze')
    logger = context['logger']
    
    try:
        # Java 소스 폴더 기본값 설정
        if not java_source_folder:
            java_source_folder = os.getenv("JAVA_SOURCE_FOLDER", ".")
            logger.info(f"Using default Java source folder: {java_source_folder}")
        
        # 옵션 검증
        _validate_analyze_options(db_object, java_object, class_name, update, java_source_folder)
        
        # 프로젝트명 결정
        detected_project_name = extract_project_name(java_source_folder)
        final_project_name = _get_or_determine_project_name(project_name, detected_project_name, java_source_folder, logger)
        
        # 분석 실행
        result = analyze_project(
            java_source_folder=java_source_folder,
            project_name=final_project_name,
            db_script_folder=db_script_folder,
            neo4j_uri=neo4j_uri,
            neo4j_user=neo4j_user,
            neo4j_password=neo4j_password,
            neo4j_database=neo4j_database,
            clean=clean,
            dry_run=dry_run,
            concurrent=concurrent,
            workers=workers,
            java_object=java_object,
            db_object=db_object,
            all_objects=all_objects,
            class_name=class_name,
            update=update,
            logger=logger
        )
        
        # end() 함수로 공통 정리
        end(context, result)
        
    except Exception as e:
        logger.error(f"Analysis failed: {e}")
        end(context, {'success': False, 'error': str(e)})
        raise click.ClickException(f"Analysis failed: {e}")

def analyze_project(java_source_folder, project_name, db_script_folder, neo4j_uri, neo4j_user, neo4j_password, neo4j_database, clean, dry_run, concurrent, workers, java_object, db_object, all_objects, class_name, update, logger):
    """실제 분석 로직을 수행하는 함수"""
    overall_start_time = datetime.now()
    java_stats = None
    db_stats = None
    
    # DB script folder가 제공되지 않으면 환경변수에서 읽기
    if not db_script_folder:
        db_script_folder = os.getenv("DB_SCRIPT_FOLDER")
        if db_script_folder:
            logger.info(f"Using DB_SCRIPT_FOLDER from environment: {db_script_folder}")
    
    try:
        # Neo4j 연결 설정
        if not dry_run:
            if not neo4j_password:
                raise ValueError("NEO4J_PASSWORD is required for database operations")
            
            db = GraphDB(neo4j_uri, neo4j_user, neo4j_password, neo4j_database)
            
            # 데이터베이스 정리
            if clean:
                if all_objects:
                    logger.info("Cleaning all database objects...")
                    db.clean_database()
                    logger.info("All database objects cleaned successfully")
                elif java_object:
                    logger.info("Cleaning Java objects only...")
                    db.clean_java_objects()
                    logger.info("Java objects cleaned successfully")
                elif db_object:
                    logger.info("Cleaning DB objects only...")
                    db.clean_db_objects()
                    logger.info("DB objects cleaned successfully")
        else:
            db = None
            logger.info("Running in dry-run mode - no database operations will be performed")
        
        # 분석 실행
        if all_objects or java_object:
            # Java 파싱
            packages_to_add, classes_to_add, class_to_package_map, beans, dependencies, endpoints, mybatis_mappers, jpa_entities, jpa_repositories, jpa_queries, config_files, test_classes, sql_statements, detected_project_name = _analyze_full_project_java(java_source_folder, project_name, logger)
            
            # Java 객체들을 Neo4j에 저장하고 통계 얻기
            java_stats = _save_java_objects_to_neo4j(db, packages_to_add, classes_to_add, class_to_package_map, beans, dependencies, endpoints, mybatis_mappers, jpa_entities, jpa_repositories, jpa_queries, config_files, test_classes, sql_statements, project_name, clean, concurrent, workers, logger)
        
        if all_objects or db_object:
            if db_script_folder:
                # DB script folder 존재 여부 확인
                if os.path.exists(db_script_folder):
                    db_stats = _analyze_full_project_db(db, db_script_folder, project_name, dry_run, logger)
                else:
                    logger.warning(f"Database script folder does not exist: {db_script_folder}")
                    logger.info("Please check the DB_SCRIPT_FOLDER path in your .env file or use --db-script-folder option")
            else:
                logger.warning("Database script folder not provided - skipping database analysis")
                logger.info("To analyze database objects, use --db-script-folder option to specify the path to SQL script files")
        
        overall_end_time = datetime.now()
        
        # Summary 출력
        print_analysis_summary(overall_start_time, overall_end_time, java_stats, db_stats, dry_run)
        
        return {
            'success': True,
            'message': 'Analysis completed successfully',
            'stats': {
                'java_stats': java_stats,
                'db_stats': db_stats
            }
        }
        
    except Exception as e:
        logger.error(f"Analysis error: {e}")
        return {
            'success': False,
            'error': str(e)
        }
    finally:
        if db:
            db.close()

if __name__ == '__main__':
    try:
        cli()
    finally:
        # 애플리케이션 종료 시 Connection Pool 정리
        pool = get_connection_pool()
        if pool.is_initialized():
            pool.close_all()
